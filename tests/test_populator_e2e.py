"""End-to-end populator test using synthetic data via monkeypatching the
yfinance + FRED clients. Verifies the full populate_ticker flow produces a
workbook with filled cells across all sheets."""
from __future__ import annotations

import shutil

import pandas as pd
import pytest
from openpyxl import load_workbook

from xl import build_workbooks, populator


def _q_columns(n=8):
    return pd.date_range("2023-12-31", periods=n, freq="QE")


def _stmt(rows: dict[str, list[float]], n=8) -> pd.DataFrame:
    cols = _q_columns(n)
    df = pd.DataFrame(rows, index=cols).T
    return df.reset_index().rename(columns={"index": "line"})


@pytest.fixture
def synthetic_data(monkeypatch):
    """Patch yfinance_client + fred_client to return deterministic fake data."""

    def fake_prices(ticker, period="10y", interval="1d", force=False):
        dates = pd.date_range("2024-01-01", periods=300, freq="D")
        closes = [100 + i * 0.1 + (i % 7) * 0.5 for i in range(300)]
        return pd.DataFrame({"date": dates, "Close": closes, "ticker": ticker})

    def fake_info(ticker, force=False):
        return {
            "shortName": f"{ticker} Corp",
            "longName": f"{ticker} Corporation",
            "sector": "Technology",
            "industry": "Semiconductors",
            "sharesOutstanding": 1_000_000_000,
            "marketCap": 100_000_000_000,
            "dividendYield": 0.005,
            "fiftyTwoWeekHigh": 150.0,
            "fiftyTwoWeekLow": 80.0,
        }

    is_q = _stmt({
        "Total Revenue": [10_000, 11_000, 12_000, 13_000, 14_000, 15_000, 16_000, 17_000],
        "Cost Of Revenue": [6_000, 6_500, 7_000, 7_500, 8_000, 8_500, 9_000, 9_500],
        "Gross Profit": [4_000, 4_500, 5_000, 5_500, 6_000, 6_500, 7_000, 7_500],
        "Selling General And Administration": [1_000, 1_100, 1_200, 1_300, 1_400, 1_500, 1_600, 1_700],
        "Research And Development": [800, 850, 900, 950, 1_000, 1_050, 1_100, 1_150],
        "Reconciled Depreciation": [200, 220, 240, 260, 280, 300, 320, 340],
        "Operating Income": [2_000, 2_330, 2_660, 2_990, 3_320, 3_650, 3_980, 4_310],
        "Interest Expense": [100, 100, 100, 100, 100, 100, 100, 100],
        "Pretax Income": [1_900, 2_230, 2_560, 2_890, 3_220, 3_550, 3_880, 4_210],
        "Tax Provision": [380, 446, 512, 578, 644, 710, 776, 842],
        "Net Income": [1_520, 1_784, 2_048, 2_312, 2_576, 2_840, 3_104, 3_368],
    })
    bs_q = _stmt({
        "Cash And Cash Equivalents": [5_000, 5_200, 5_500, 5_800, 6_100, 6_400, 6_700, 7_000],
        "Accounts Receivable": [2_000, 2_100, 2_200, 2_300, 2_400, 2_500, 2_600, 2_700],
        "Inventory": [1_500, 1_550, 1_600, 1_650, 1_700, 1_750, 1_800, 1_850],
        "Current Assets": [10_000, 10_300, 10_600, 10_900, 11_200, 11_500, 11_800, 12_100],
        "Total Assets": [50_000, 51_000, 52_000, 53_000, 54_000, 55_000, 56_000, 57_000],
        "Current Liabilities": [5_000, 5_100, 5_200, 5_300, 5_400, 5_500, 5_600, 5_700],
        "Long Term Debt": [10_000, 10_000, 10_000, 10_000, 10_000, 10_000, 10_000, 10_000],
        "Total Liabilities Net Minority Interest": [20_000, 20_100, 20_200, 20_300, 20_400, 20_500, 20_600, 20_700],
        "Retained Earnings": [15_000, 16_000, 17_000, 18_000, 19_000, 20_000, 21_000, 22_000],
        "Common Stock Equity": [30_000, 30_900, 31_800, 32_700, 33_600, 34_500, 35_400, 36_300],
    })
    cf_q = _stmt({
        "Operating Cash Flow": [2_500, 2_800, 3_100, 3_400, 3_700, 4_000, 4_300, 4_600],
        "Capital Expenditure": [-500, -550, -600, -650, -700, -750, -800, -850],
        "Free Cash Flow": [2_000, 2_250, 2_500, 2_750, 3_000, 3_250, 3_500, 3_750],
    })

    def fake_is(ticker, quarterly=True, force=False):
        return is_q

    def fake_bs(ticker, quarterly=True, force=False):
        return bs_q

    def fake_cf(ticker, quarterly=True, force=False):
        return cf_q

    def fake_eps_est(ticker, force=False):
        return pd.DataFrame({
            "period": ["0y", "+1y"],
            "avg": [12.0, 14.5],
        })

    def fake_rev_est(ticker, force=False):
        return pd.DataFrame({
            "period": ["0y", "+1y"],
            "avg": [62_000, 75_000],
        })

    def fake_revisions(ticker, force=False):
        return pd.DataFrame({"upLast30days": [3], "upLast7days": [1]})

    def fake_calendar(ticker, force=False):
        return {"Earnings Date": "2026-07-28"}

    def fake_insider(ticker, force=False):
        return pd.DataFrame()

    def fake_iv(ticker, force=False):
        return 0.45

    def fake_hv(ticker, window_days=30):
        return 0.35

    def fake_rf():
        return 0.042

    def fake_curve():
        return pd.DataFrame([
            {"series": "DGS3MO", "label": "3M", "yield_pct": 4.2},
            {"series": "DGS10", "label": "10Y", "yield_pct": 4.2},
        ])

    def fake_beta(ticker, market="SPY", months=60):
        return 1.3

    import compute.capm as capm_mod
    from data import fred_client, yfinance_client

    monkeypatch.setattr(yfinance_client, "prices", fake_prices)
    monkeypatch.setattr(yfinance_client, "info", fake_info)
    monkeypatch.setattr(yfinance_client, "income_statement", fake_is)
    monkeypatch.setattr(yfinance_client, "balance_sheet", fake_bs)
    monkeypatch.setattr(yfinance_client, "cashflow", fake_cf)
    monkeypatch.setattr(yfinance_client, "earnings_estimate", fake_eps_est)
    monkeypatch.setattr(yfinance_client, "revenue_estimate", fake_rev_est)
    monkeypatch.setattr(yfinance_client, "eps_revisions", fake_revisions)
    monkeypatch.setattr(yfinance_client, "calendar", fake_calendar)
    monkeypatch.setattr(yfinance_client, "insider_transactions", fake_insider)
    monkeypatch.setattr(yfinance_client, "options_atm_iv", fake_iv)
    monkeypatch.setattr(yfinance_client, "historical_vol", fake_hv)
    monkeypatch.setattr(fred_client, "risk_free_10y", fake_rf)
    monkeypatch.setattr(fred_client, "treasury_curve", fake_curve)
    monkeypatch.setattr(capm_mod, "beta", fake_beta)

    # populator imports beta as part of `capm` module attribute reference
    monkeypatch.setattr(populator.capm, "beta", fake_beta)
    monkeypatch.setattr(populator.yfc, "prices", fake_prices)
    monkeypatch.setattr(populator.yfc, "info", fake_info)
    monkeypatch.setattr(populator.yfc, "income_statement", fake_is)
    monkeypatch.setattr(populator.yfc, "balance_sheet", fake_bs)
    monkeypatch.setattr(populator.yfc, "cashflow", fake_cf)
    monkeypatch.setattr(populator.yfc, "earnings_estimate", fake_eps_est)
    monkeypatch.setattr(populator.yfc, "revenue_estimate", fake_rev_est)
    monkeypatch.setattr(populator.yfc, "eps_revisions", fake_revisions)
    monkeypatch.setattr(populator.yfc, "calendar", fake_calendar)
    monkeypatch.setattr(populator.yfc, "insider_transactions", fake_insider)
    monkeypatch.setattr(populator.yfc, "options_atm_iv", fake_iv)
    monkeypatch.setattr(populator.yfc, "historical_vol", fake_hv)
    monkeypatch.setattr(populator.fc, "risk_free_10y", fake_rf)


def test_populate_ticker_end_to_end(synthetic_data, tmp_path):
    template = tmp_path / "template.xlsx"
    build_workbooks.build_ticker_template(template)
    out = tmp_path / "Ticker_TEST.xlsx"

    result = populator.populate_ticker(template, out, "TEST")
    assert result.exists()

    wb = load_workbook(result)

    # Cover
    cover = wb["Cover"]
    ticker_row = populator.find_label_row(cover, "Ticker")
    assert cover.cell(row=ticker_row, column=2).value == "TEST"
    price_row = populator.find_label_row(cover, "Price")
    assert cover.cell(row=price_row, column=2).value is not None

    # Market — CAPM labels live in col 4, values in col 5
    mkt = wb["Market"]
    re_cell = populator.find_label_cell(mkt, "Cost of Equity (Re)")
    assert re_cell is not None
    re_row, re_col = re_cell
    re = mkt.cell(row=re_row, column=re_col + 1).value
    assert re is not None and 0 < re < 0.5

    # Fin Stat
    fs = wb["Fin Stat"]
    rev_row = populator.find_label_row(fs, "Revenue")
    # 8 quarters of data populated across cols 2..9
    populated_cols = [c for c in range(2, 14)
                      if fs.cell(row=rev_row, column=c).value is not None]
    assert len(populated_cols) >= 4

    # Analysis — per-quarter values across cols 2..9 (8 quarters of fake data)
    an = wb["Analysis"]
    cr_row = populator.find_label_row(an, "Current Ratio (Liquidity)")
    populated = [an.cell(row=cr_row, column=c).value for c in range(2, 10)
                 if an.cell(row=cr_row, column=c).value is not None]
    assert len(populated) >= 4, f"expected per-quarter ratios, got {len(populated)}"
    cr_score = an.cell(row=cr_row, column=13).value
    assert isinstance(cr_score, int)

    # Spot-check that newest quarter ratio differs from oldest (data is rising)
    op_row = populator.find_label_row(an, "Operating Margin")
    om_oldest = an.cell(row=op_row, column=2).value
    om_newest = an.cell(row=op_row, column=9).value
    assert om_oldest is not None and om_newest is not None
    assert om_newest > om_oldest  # margins rising in fake data

    # Valuation
    val = wb["Valuation"]
    rev_forecast_row = populator.find_label_row(val, "Revenue")
    assert val.cell(row=rev_forecast_row, column=2).value is not None  # Y+1

    # Credit (Z-Score lives in col 4; CAPM/Altman labels are right-block)
    cr = wb["Credit"]
    edf_row = populator.find_label_row(cr, "EDF (Expected Default Frequency)")
    edf = cr.cell(row=edf_row, column=2).value
    assert edf is not None and 0 <= edf <= 1

    z_cell = populator.find_label_cell(cr, "Z-Score")
    assert z_cell is not None
    z = cr.cell(row=z_cell[0], column=z_cell[1] + 1).value
    assert z is not None

    # Multiples band
    pe_row = populator.find_label_row(an, "P/E")
    assert an.cell(row=pe_row, column=2).value is not None  # current P/E filled

    # Owner Earnings
    oe_row = populator.find_label_row(an, "Owner Earnings (TTM)")
    assert an.cell(row=oe_row, column=2).value is not None
    cap_row = populator.find_label_row(an, "Capital Allocation: CAPEX %")
    assert an.cell(row=cap_row, column=2).value is not None


def test_populate_positions_with_correlations(tmp_path, monkeypatch):
    template = tmp_path / "Market_Daily.xlsx"
    build_workbooks.build_market_daily(template)

    # Pre-fill positions: 3 tickers with shares + cost
    wb = load_workbook(template)
    pos = wb["Positions"]
    pos.cell(row=4, column=1, value="AAA")
    pos.cell(row=4, column=2, value="Alpha Co")
    pos.cell(row=4, column=3, value=100)
    pos.cell(row=4, column=4, value=50.0)
    pos.cell(row=5, column=1, value="BBB")
    pos.cell(row=5, column=3, value=200)
    pos.cell(row=5, column=4, value=30.0)
    pos.cell(row=6, column=1, value="CCC")
    pos.cell(row=6, column=3, value=50)
    pos.cell(row=6, column=4, value=80.0)
    wb.save(template)

    # Patch prices: AAA & BBB share the same daily returns; CCC uses inverted returns
    import numpy as np
    rng = np.random.default_rng(seed=42)
    base_returns = rng.normal(0.0005, 0.015, 120)

    def make_path(returns, start):
        path = [start]
        for r in returns:
            path.append(path[-1] * (1 + r))
        return path[1:]

    def fake_prices(ticker, period="1y", interval="1d", force=False):
        dates = pd.date_range("2024-01-01", periods=120, freq="D")
        if ticker == "AAA":
            closes = make_path(base_returns, 100)
        elif ticker == "BBB":
            closes = make_path(base_returns, 50)
        elif ticker == "CCC":
            closes = make_path(-base_returns, 200)
        else:
            closes = [100] * 120
        return pd.DataFrame({"date": dates, "Close": closes, "ticker": ticker})

    monkeypatch.setattr(populator.yfc, "prices", fake_prices)
    monkeypatch.setattr(populator.yfc, "calendar", lambda t, force=False: {})
    populator.populate_market_daily(template)

    wb2 = load_workbook(template)
    pos2 = wb2["Positions"]

    assert pos2.cell(row=4, column=5).value is not None
    # Correlation matrix anchor is at row 33 (header), matrix starts row 35
    aaa_bbb_corr = pos2.cell(row=35, column=3).value  # AAA row, BBB col
    assert aaa_bbb_corr is not None
    assert abs(aaa_bbb_corr - 1.0) < 0.01  # same returns → corr ~1

    # AAA & CCC have inverted daily returns → corr ~ -1
    aaa_ccc_corr = pos2.cell(row=35, column=4).value  # AAA row, CCC col
    assert aaa_ccc_corr is not None
    assert abs(aaa_ccc_corr - (-1.0)) < 0.01


def test_screener_flags_double_buy_signal(tmp_path, monkeypatch):
    template = tmp_path / "Market_Daily.xlsx"
    build_workbooks.build_market_daily(template)

    wb = load_workbook(template)
    pos = wb["Positions"]
    pos.cell(row=4, column=1, value="UP")
    pos.cell(row=4, column=2, value="Uptrend Co")
    pos.cell(row=5, column=1, value="DOWN")
    pos.cell(row=5, column=2, value="Downtrend Co")
    wb.save(template)

    def fake_prices(ticker, period="2y", interval="1d", force=False):
        dates = pd.date_range("2024-01-01", periods=300, freq="D")
        if ticker == "UP":
            closes = [100 + i * 0.5 for i in range(300)]   # clear uptrend
        elif ticker == "DOWN":
            closes = [200 - i * 0.5 for i in range(300)]   # clear downtrend
        else:
            closes = [100] * 300
        return pd.DataFrame({"date": dates, "Close": closes, "ticker": ticker})

    monkeypatch.setattr(populator.yfc, "prices", fake_prices)
    monkeypatch.setattr(populator.yfc, "calendar", lambda t, force=False: {})
    populator.populate_market_daily(template)

    wb2 = load_workbook(template)
    scr = wb2["Screener"]
    # Row 4 should be the UP ticker (DOWN should not flag)
    assert scr.cell(row=4, column=1).value == "UP"
    assert scr.cell(row=4, column=3).value == "YES"
    assert scr.cell(row=4, column=4).value == "YES"
    # Row 5 should be empty (DOWN doesn't pass)
    assert scr.cell(row=5, column=1).value is None


def test_summary_aggregates_ticker_files(tmp_path, monkeypatch):
    """When Tickers/ contains multiple populated ticker workbooks, the
    Summary tab on Market_Daily should aggregate them sorted by upside."""
    import config as cfg
    monkeypatch.setattr(cfg, "_SETTINGS", None)

    # Build two minimal ticker files with different Cover values
    template = tmp_path / "Ticker_TEMPLATE.xlsx"
    build_workbooks.build_ticker_template(template)

    def _make_ticker(out_path, ticker, upside, edf, fair_value, quality):
        import shutil
        shutil.copy(template, out_path)
        wb = load_workbook(out_path)
        cover = wb["Cover"]
        populator.write_label_value(cover, "Ticker", ticker)
        populator.write_label_value(cover, "Name", f"{ticker} Corp")
        populator.write_label_value(cover, "Sector", "Tech")
        populator.write_label_value(cover, "Price", 100.0)
        populator.write_label_value(cover, "Fair Value (DCF)", fair_value)
        populator.write_label_value(cover, "Upside %", upside)
        populator.write_label_value(cover, "Margin of Safety", 0.2)
        populator.write_label_value(cover, "Quality Score (avg)", quality)
        populator.write_label_value(cover, "EDF 1y (KMV)", edf)
        populator.write_label_value(cover, "Altman Z + Bucket", "2.5 (Grey)")
        wb.save(out_path)

    tickers_dir = tmp_path / "Tickers"
    tickers_dir.mkdir()
    _make_ticker(tickers_dir / "Ticker_HIGH.xlsx", "HIGH", 0.60, 0.02, 160.0, 8.5)
    _make_ticker(tickers_dir / "Ticker_LOW.xlsx", "LOW", 0.10, 0.08, 110.0, 5.0)
    _make_ticker(tickers_dir / "Ticker_MID.xlsx", "MID", 0.30, 0.04, 130.0, 7.0)

    # Point config at tmp dir
    monkeypatch.setattr(cfg, "load", lambda: {
        "tickers_dir": tickers_dir,
        "cache_dir": tmp_path / "Cache",
        "archive_dir": tmp_path / "Archive",
        "universe": {},
        "cache_ttl_days": {},
        "discount_rates": [0.04],
        "terminal_multiples": [10.0],
    })

    market_path = tmp_path / "Market_Daily.xlsx"
    build_workbooks.build_market_daily(market_path)

    monkeypatch.setattr(populator.yfc, "prices",
                        lambda t, period="2y", interval="1d", force=False: pd.DataFrame())
    monkeypatch.setattr(populator.yfc, "calendar", lambda t, force=False: {})

    populator.populate_market_daily(market_path)

    wb = load_workbook(market_path)
    summary = wb["Summary"]
    # Sorted by upside descending: HIGH, MID, LOW
    assert summary.cell(row=4, column=1).value == "HIGH"
    assert summary.cell(row=5, column=1).value == "MID"
    assert summary.cell(row=6, column=1).value == "LOW"
    # Quality scores match
    assert summary.cell(row=4, column=8).value == pytest.approx(8.5)
    # EDF column populated
    assert summary.cell(row=4, column=9).value == pytest.approx(0.02)


def test_snapshot_to_archive(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    src = tmp_path / "Ticker_FOO.xlsx"
    wb.save(src)
    archive = tmp_path / "Archive"

    snap = populator.snapshot_to_archive(src, archive_dir=archive)
    assert snap.exists()
    assert snap.parent == archive
    assert snap.stem.startswith("Ticker_FOO_")


def test_populate_ticker_asset_light_hides_dso(synthetic_data, tmp_path,
                                                monkeypatch):
    # Override industry to "Software" → asset_light = True
    def software_info(ticker, force=False):
        return {
            "shortName": "Software Co",
            "industry": "Software—Application",
            "sector": "Technology",
            "sharesOutstanding": 1_000_000_000,
            "marketCap": 100_000_000_000,
            "dividendYield": 0,
            "fiftyTwoWeekHigh": 150.0,
            "fiftyTwoWeekLow": 80.0,
        }
    monkeypatch.setattr(populator.yfc, "info", software_info)

    template = tmp_path / "template.xlsx"
    build_workbooks.build_ticker_template(template)
    out = tmp_path / "Ticker_SOFT.xlsx"
    populator.populate_ticker(template, out, "SOFT")

    wb = load_workbook(out)
    an = wb["Analysis"]
    dso_row = populator.find_label_row(an, "DSO (Efficiency)")
    inv_row = populator.find_label_row(an, "Inventory Turnover (Efficiency)")
    assert an.cell(row=dso_row, column=2).value is None
    assert an.cell(row=dso_row, column=13).value == "N/A"
    assert an.cell(row=inv_row, column=2).value is None
