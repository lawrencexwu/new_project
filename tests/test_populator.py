"""Tests for the populator. Uses synthetic dataframes that match the yfinance
shape (line items as index, dates as columns)."""
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from xl import build_workbooks, populator


@pytest.fixture(scope="module")
def template_path(tmp_path_factory):
    out = tmp_path_factory.mktemp("wb") / "Ticker_TEMPLATE.xlsx"
    build_workbooks.build_ticker_template(out)
    return out


@pytest.fixture
def workbook(template_path, tmp_path):
    path = tmp_path / "Ticker_TEST.xlsx"
    import shutil
    shutil.copy(template_path, path)
    return load_workbook(path), path


def _q_columns(n=4):
    return pd.date_range("2024-03-31", periods=n, freq="QE")


def _stmt(rows: dict[str, list[float]], n=4) -> pd.DataFrame:
    cols = _q_columns(n)
    df = pd.DataFrame(rows, index=cols).T
    df = df.reset_index().rename(columns={"index": "line"})
    return df


def test_find_label_row(workbook):
    wb, _ = workbook
    ws = wb["Fin Stat"]
    row = populator.find_label_row(ws, "Revenue")
    assert row is not None and row > 0


def test_find_section_anchor(workbook):
    wb, _ = workbook
    ws = wb["Daily Plan"] if "Daily Plan" in wb.sheetnames else wb["Analysis"]
    # Analysis has section "CFO / Earnings Quality"
    anchor = populator.find_section_anchor(wb["Analysis"], "CFO / Earnings Quality")
    assert anchor is not None


def test_yf_line_mapping():
    df = _stmt({"Total Revenue": [100, 110, 120, 130]})
    row = populator._row_from_yf(df, "Revenue")
    assert row is not None
    assert row.iloc[-1] == 130


def test_yf_line_mapping_alt_label():
    df = _stmt({"Operating Revenue": [50, 60, 70, 80]})
    row = populator._row_from_yf(df, "Revenue")
    assert row is not None and row.iloc[-1] == 80


def test_yf_line_mapping_missing():
    df = _stmt({"Other Stuff": [1, 2, 3, 4]})
    row = populator._row_from_yf(df, "Revenue")
    assert row is None


def test_ttm_sum():
    df = _stmt({"Total Revenue": [100, 110, 120, 130]})
    assert populator._ttm(df, "Revenue") == 460


def test_ttm_uses_newest_quarters_even_when_columns_reversed():
    """yfinance returns columns in descending date order. iloc[-4:] would
    give the OLDEST 4 quarters without sorting — that bug made NVDA's
    Owner Earnings TTM ~$99B instead of ~$120B."""
    dates_desc = pd.date_range("2025-06-30", periods=5, freq="-1QE")
    df = pd.DataFrame({"Total Revenue": [50, 40, 30, 20, 10]}, index=dates_desc).T
    df = df.reset_index().rename(columns={"index": "line"})
    # Last 4 IN DATE ORDER are 20, 30, 40, 50 → 140
    # If we naively iloc[-4:] on the input columns we'd get 40, 30, 20, 10 → 100
    assert populator._ttm(df, "Revenue") == 140


def test_latest_value_uses_newest_quarter():
    dates_desc = pd.date_range("2025-06-30", periods=4, freq="-1QE")
    df = pd.DataFrame({"Total Revenue": [50, 40, 30, 20]}, index=dates_desc).T
    df = df.reset_index().rename(columns={"index": "line"})
    # Newest (2025-06-30) value is 50
    assert populator._latest_value(df, "Revenue") == 50


def test_format_earnings_date():
    import datetime as dt
    assert populator._format_earnings_date([dt.date(2026, 5, 21)]) == "2026-05-21"
    assert populator._format_earnings_date(dt.date(2026, 7, 28)) == "2026-07-28"
    assert populator._format_earnings_date("2026-08-15") == "2026-08-15"
    assert populator._format_earnings_date(None) == ""
    assert populator._format_earnings_date([]) == ""


def test_fade_growth():
    g = populator._fade_growth(0.20, 0.15, years=5, terminal=0.025)
    assert len(g) == 5
    assert g[0] == 0.20
    assert g[1] == 0.15
    assert g[-1] < g[1]  # fading toward terminal


def test_safe_helpers():
    assert populator._safe_float("12.5") == 12.5
    assert populator._safe_float(None) is None
    assert populator._safe_float("abc") is None
    assert populator._safe_sum(1, 2, 3) == 6
    assert populator._safe_sum(None, None) is None
    assert populator._safe_sum(5, None) == 5


def test_consensus_growth_defaults_when_empty():
    g1, g2 = populator._consensus_growth(pd.DataFrame(), 1000)
    assert g1 == 0.05
    assert g2 == 0.04


def test_populate_fin_stat_writes_revenue(workbook):
    wb, path = workbook
    is_q = _stmt({"Total Revenue": [100, 110, 120, 130],
                  "Cost Of Revenue": [60, 65, 70, 75],
                  "Operating Income": [10, 12, 14, 16]})
    bs_q = _stmt({"Total Assets": [500, 510, 520, 530],
                  "Current Liabilities": [100, 105, 110, 115],
                  "Current Assets": [200, 210, 220, 230]})
    cf_q = _stmt({"Operating Cash Flow": [20, 22, 24, 26],
                  "Capital Expenditure": [-5, -6, -7, -8]})
    populator._populate_fin_stat(wb, is_q, bs_q, cf_q)
    wb.save(path)
    wb2 = load_workbook(path)
    ws = wb2["Fin Stat"]
    rev_row = populator.find_label_row(ws, "Revenue")
    # Latest value is in column 5 (4 quarters starting at col 2)
    assert ws.cell(row=rev_row, column=5).value == 130


def test_score_band_map_lookups():
    from compute import scoring
    assert scoring.score(0.8, populator._SCORE_BAND_MAP["Quality of Earnings (OCF/NI)"]) is not None
    assert scoring.score(1.5, populator._SCORE_BAND_MAP["Current Ratio (Liquidity)"]) == 9
    assert scoring.score(5.0, populator._SCORE_BAND_MAP["Interest Cover"]) == 7
    assert scoring.score(44.0, populator._SCORE_BAND_MAP["DSO (Efficiency)"]) == 9


def test_red_flags_formatting():
    scores = {"Critical Debt Load": 3, "Interest Cover": 4, "Quality": 7}
    out = populator._format_red_flags(scores)
    assert len(out) == 2
    assert "Critical Debt Load" in out[0]  # worst first


def test_template_has_analysis_charts(template_path):
    wb = load_workbook(template_path)
    ws = wb["Analysis"]
    assert len(ws._charts) == 4, f"expected 4 charts on Analysis, got {len(ws._charts)}"


def test_asset_light_detection():
    # Simulating populate_ticker's gate decision via the industry constants
    industries = [
        ("Software—Application", True),
        ("Insurance—Diversified", True),
        ("Investment Banking", True),
        ("Auto Manufacturers", False),
        ("Steel", False),
        ("Internet Content & Information", True),
    ]
    for industry, expected in industries:
        is_light = any(h in industry.lower() for h in populator.ASSET_LIGHT_INDUSTRY_HINTS)
        assert is_light == expected, f"{industry} expected {expected}"
