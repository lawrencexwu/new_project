"""Populate workbook cells with live data. Uses openpyxl so it runs headless.

The xlwings macros in xl/macros.py are thin wrappers that call these functions.
"""
from __future__ import annotations

import math
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_BOLD_FONT = Font(name="Calibri", size=10, bold=True)
_SPARKLINE_FONT = Font(name="Consolas", size=11)
_BLOCK_CHARS = "▁▂▃▄▅▆▇█"


def block_sparkline(values, width: int = 20) -> str:
    """Render a list of numbers as a Unicode-block sparkline string.
    Adapts to `width` characters; auto-scales between min and max."""
    if values is None:
        return ""
    vals = [float(v) for v in values
            if v is not None and not (isinstance(v, float) and math.isnan(v))]
    if len(vals) < 2:
        return ""
    # Resample if more values than width
    if len(vals) > width:
        step = len(vals) / width
        sampled = [vals[int(i * step)] for i in range(width)]
    else:
        sampled = vals
    lo, hi = min(sampled), max(sampled)
    if hi == lo:
        return _BLOCK_CHARS[3] * len(sampled)
    out = []
    for v in sampled:
        idx = int((v - lo) / (hi - lo) * (len(_BLOCK_CHARS) - 1))
        out.append(_BLOCK_CHARS[idx])
    return "".join(out)

import config
from compute import altman, capm, dcf, hillegeist, kmv, merton, multiples, ratios, scoring, signals
from data import edgar_client as edgar
from data import fred_client as fc
from data import yfinance_client as yfc
from xl import sparkline_injector


ASSET_LIGHT_INDUSTRY_HINTS = (
    "software", "saas", "internet", "advertising", "marketing", "consulting",
    "media", "broadcast", "publishing", "financial data", "asset management",
    "investment", "broker", "exchange", "insurance", "bank",
)


# -----------------------------------------------------------
# Generic helpers
# -----------------------------------------------------------

_DEFAULT_LABEL_COLS = (1, 4, 7, 9, 10)


def find_label_row(ws: Worksheet, label: str, col: int | None = None,
                   max_row: int = 300) -> int | None:
    """Find row of an exact label match. When col is None, search the standard
    set of label columns (1, 4, 7, 9, 10) and return on first hit.
    Returns the row only; callers wanting the column should use find_label_cell."""
    cell = find_label_cell(ws, label, col=col, max_row=max_row)
    return cell[0] if cell else None


def find_label_cell(ws: Worksheet, label: str, col: int | None = None,
                    max_row: int = 300) -> tuple[int, int] | None:
    """Return (row, col) of the first cell whose value equals label."""
    cols = (col,) if col is not None else _DEFAULT_LABEL_COLS
    for c in cols:
        for row in range(1, max_row + 1):
            v = ws.cell(row=row, column=c).value
            if v == label:
                return (row, c)
    return None


def find_section_anchor(ws: Worksheet, section_label: str,
                        cols: tuple[int, ...] = (1, 4, 9)) -> int | None:
    for col in cols:
        r = find_label_row(ws, section_label, col=col)
        if r is not None:
            return r
    return None


def set_value(ws: Worksheet, row: int, col: int, value: Any) -> None:
    if row is None or col is None:
        return
    cell = ws.cell(row=row, column=col)
    if value is None or (isinstance(value, float) and math.isnan(value)):
        cell.value = None
        return
    cell.value = value


def write_label_value(ws: Worksheet, label: str, value: Any,
                      label_col: int | None = None,
                      value_col_offset: int = 1) -> None:
    """Write `value` next to a label. Searches all standard label columns
    when label_col is None — supports left-block and side-block layouts."""
    found = find_label_cell(ws, label, col=label_col)
    if found is None:
        return
    row, col = found
    set_value(ws, row, col + value_col_offset, value)


# -----------------------------------------------------------
# yfinance line-item label mapping
# -----------------------------------------------------------

_YF_LINE_MAP = {
    # Income statement
    "Revenue": ("Total Revenue", "Operating Revenue", "TotalRevenue"),
    "Cost of Revenue": ("Cost Of Revenue", "Reconciled Cost Of Revenue"),
    "Gross Profit": ("Gross Profit",),
    "SG&A": ("Selling General And Administration", "Selling General And Administrative"),
    "R&D": ("Research And Development",),
    "Depreciation & Amortization": ("Reconciled Depreciation", "Depreciation And Amortization"),
    "Other Opex": ("Other Operating Expenses",),
    "Operating Income (EBIT)": ("Operating Income", "EBIT"),
    "Interest Expense": ("Interest Expense", "Interest Expense Non Operating"),
    "Other Non-Op Income": ("Other Income Expense", "Non Operating Income"),
    "Pretax Income": ("Pretax Income", "Income Before Tax"),
    "Provision for Taxes": ("Tax Provision", "Income Tax Expense"),
    "Net Income": ("Net Income", "Net Income Common Stockholders"),
    # Balance sheet
    "Cash & ST Investments": ("Cash Cash Equivalents And Short Term Investments",
                              "Cash And Cash Equivalents"),
    "Receivables": ("Accounts Receivable", "Receivables"),
    "Inventory": ("Inventory",),
    "Current Assets": ("Current Assets", "Total Current Assets"),
    "Total Assets": ("Total Assets",),
    "Current Liabilities": ("Current Liabilities", "Total Current Liabilities"),
    "Long-Term Debt": ("Long Term Debt", "Long Term Debt And Capital Lease Obligation"),
    "Total Liabilities": ("Total Liabilities Net Minority Interest", "Total Liab"),
    "Retained Earnings": ("Retained Earnings",),
    "Total Equity": ("Common Stock Equity", "Total Equity Gross Minority Interest"),
    "Minority Interest": ("Minority Interest",),
    # Cash flow
    "Operating Cash Flow": ("Operating Cash Flow", "Cash Flow From Continuing Operating Activities"),
    "Capital Expenditures": ("Capital Expenditure",),
    "Free Cash Flow": ("Free Cash Flow",),
    "Dividends Paid": ("Cash Dividends Paid", "Common Stock Dividend Paid"),
    "Share Buybacks": ("Repurchase Of Capital Stock", "Common Stock Issuance"),
    "Net Debt Issued/Repaid": ("Net Issuance Payments Of Debt",),
}


def _row_from_yf(df: pd.DataFrame, label: str) -> pd.Series | None:
    """yfinance returns line items as index, dates as columns. Find the row."""
    if df is None or df.empty:
        return None
    if "line" in df.columns:
        index = df["line"].astype(str)
    else:
        index = pd.Series(df.index.astype(str), index=df.index)
    candidates = _YF_LINE_MAP.get(label, (label,))
    for c in candidates:
        matches = index[index.str.casefold() == c.casefold()]
        if not matches.empty:
            i = matches.index[0]
            if "line" in df.columns:
                row = df.loc[i].drop("line", errors="ignore")
            else:
                row = df.loc[i]
            return row
    return None


def _quarter_columns(df: pd.DataFrame, n: int = 11) -> list:
    """Return the last n columns (period dates) ordered oldest→newest."""
    if df is None or df.empty:
        return []
    cols = list(df.columns)
    if "line" in cols:
        cols.remove("line")
    try:
        cols_sorted = sorted(cols, key=lambda c: pd.to_datetime(c, errors="coerce") or pd.NaT)
    except TypeError:
        cols_sorted = cols
    return cols_sorted[-n:]


def _latest_value(df: pd.DataFrame, label: str) -> float | None:
    row = _row_from_yf(df, label)
    if row is None or row.empty:
        return None
    row = _sort_by_date_index(row).dropna()
    if row.empty:
        return None
    try:
        return float(row.iloc[-1])
    except (TypeError, ValueError):
        return None


def _sort_by_date_index(row: pd.Series) -> pd.Series:
    """Reindex by date so iloc[-1] is the newest column. yfinance returns
    columns in descending date order, so naive iloc[-1] would give the
    oldest quarter."""
    try:
        idx_dt = pd.to_datetime(row.index, errors="coerce")
        if idx_dt.isna().all():
            return row
        sorted_row = row.copy()
        sorted_row.index = idx_dt
        return sorted_row.sort_index()
    except Exception:
        return row


# -----------------------------------------------------------
# Ticker workbook populator
# -----------------------------------------------------------

def populate_ticker(template_path: Path, output_path: Path, ticker: str,
                    force: bool = False) -> Path:
    template_path = Path(template_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if not output_path.exists() or template_path.resolve() != output_path.resolve():
        shutil.copy(template_path, output_path)

    wb = load_workbook(output_path)
    ticker = ticker.upper()

    info = yfc.info(ticker, force=force)
    industry = (info.get("industry") or "").lower()
    asset_light = any(h in industry for h in ASSET_LIGHT_INDUSTRY_HINTS)

    px_df = yfc.prices(ticker, period="10y", force=force)
    is_q = yfc.income_statement(ticker, quarterly=True, force=force)
    bs_q = yfc.balance_sheet(ticker, quarterly=True, force=force)
    cf_q = yfc.cashflow(ticker, quarterly=True, force=force)

    # EDGAR fallback when yfinance returns empty
    if is_q is None or is_q.empty:
        is_q = edgar.income_statement(ticker, force=force)
    if bs_q is None or bs_q.empty:
        bs_q = edgar.balance_sheet(ticker, force=force)
    if cf_q is None or cf_q.empty:
        cf_q = edgar.cashflow(ticker, force=force)

    rf = fc.risk_free_10y()
    mrp = 0.05  # default market risk premium
    beta_val = capm.beta(ticker)

    market = _populate_market(wb, ticker, info, px_df, rf, beta_val, mrp)
    _populate_fin_stat(wb, is_q, bs_q, cf_q)
    analysis_results = _populate_analysis(wb, is_q, bs_q, cf_q, asset_light=asset_light)
    _populate_owner_earnings(wb, is_q, cf_q)
    _populate_multiples_band(wb, is_q, bs_q, cf_q, px_df, market)
    valuation_results = _populate_valuation(
        wb, ticker, is_q, bs_q, cf_q, px_df, market,
        force=force,
    )
    credit_results = _populate_credit(wb, market, is_q, bs_q, px_df)

    _populate_cover(
        wb,
        ticker=ticker,
        info=info,
        market=market,
        analysis=analysis_results,
        valuation=valuation_results,
        credit=credit_results,
    )

    # Auto-fit columns on every sheet so the user doesn't have to manually
    # double-click each column boundary.
    for sheet_name in wb.sheetnames:
        autofit_columns(wb[sheet_name])

    wb.save(output_path)
    return output_path


# -----------------------------------------------------------
# Section: Market & Macro
# -----------------------------------------------------------

def _populate_market(wb, ticker: str, info: dict, px_df,
                     rf: float | None, beta_val: float | None,
                     mrp: float) -> dict:
    ws = wb["Market"]
    result: dict[str, Any] = {"ticker": ticker}

    price = None
    if px_df is not None and not px_df.empty:
        price = float(px_df["Close"].iloc[-1])
        close = px_df.set_index(pd.to_datetime(px_df["date"]))["Close"].astype(float)
    else:
        close = None

    result["price"] = price
    result["close"] = close

    shares = _safe_float(info.get("sharesOutstanding"))
    div_yield = _safe_float(info.get("dividendYield"))
    if div_yield is not None and div_yield > 1.0:
        div_yield = div_yield / 100.0  # yfinance occasionally returns percentage points
    market_cap = price * shares if (price and shares) else _safe_float(info.get("marketCap"))
    fifty_two_w_high = _safe_float(info.get("fiftyTwoWeekHigh"))
    fifty_two_w_low = _safe_float(info.get("fiftyTwoWeekLow"))

    write_label_value(ws, "Price", price)
    write_label_value(ws, "Shares Outstanding", shares)
    write_label_value(ws, "Market Cap", market_cap)
    write_label_value(ws, "Dividend Yield", div_yield)
    write_label_value(ws, "52w High", fifty_two_w_high)
    write_label_value(ws, "52w Low", fifty_two_w_low)

    hv30 = yfc.historical_vol(ticker, 30)
    hv90 = yfc.historical_vol(ticker, 90)
    hv365 = yfc.historical_vol(ticker, 365)
    atm_iv = yfc.options_atm_iv(ticker)
    write_label_value(ws, "Historical Vol (30d)", hv30)
    write_label_value(ws, "Historical Vol (90d)", hv90)
    write_label_value(ws, "Historical Vol (365d)", hv365)
    write_label_value(ws, "ATM Implied Vol", atm_iv)

    write_label_value(ws, "Risk-Free Rate (10Y)", rf)
    write_label_value(ws, "Market Risk Premium", mrp)
    write_label_value(ws, "Country Risk Premium", 0.0)
    write_label_value(ws, "Beta", beta_val)

    re = capm.cost_of_equity(rf, beta_val, mrp) if (rf and beta_val) else None
    write_label_value(ws, "Cost of Equity (Re)", re)

    tax_rate = 0.21
    write_label_value(ws, "Tax Rate", tax_rate)

    result["shares"] = shares
    result["market_cap"] = market_cap
    result["rf"] = rf
    result["beta"] = beta_val
    result["mrp"] = mrp
    result["re"] = re
    result["tax_rate"] = tax_rate
    result["equity_vol"] = hv365 or hv90 or hv30
    return result


# -----------------------------------------------------------
# Section: Financial Statements
# -----------------------------------------------------------

_IS_LABELS = ["Revenue", "Cost of Revenue", "Gross Profit", "SG&A", "R&D",
              "Depreciation & Amortization", "Other Opex",
              "Operating Income (EBIT)", "Interest Expense",
              "Other Non-Op Income", "Pretax Income",
              "Provision for Taxes", "Net Income"]
_BS_LABELS = ["Cash & ST Investments", "Receivables", "Inventory",
              "Current Assets", "Total Assets",
              "Current Liabilities", "Long-Term Debt", "Total Liabilities",
              "Retained Earnings", "Total Equity", "Minority Interest"]
_CF_LABELS = ["Operating Cash Flow", "Capital Expenditures",
              "Free Cash Flow", "Dividends Paid", "Share Buybacks",
              "Net Debt Issued/Repaid"]


def _populate_fin_stat(wb, is_q, bs_q, cf_q):
    ws = wb["Fin Stat"]
    _write_statement_block(ws, _IS_LABELS, is_q, "Income Statement (Quarterly)")
    _write_statement_block(ws, _BS_LABELS, bs_q, "Balance Sheet (Quarterly)")
    _write_statement_block(ws, _CF_LABELS, cf_q, "Cash Flow (Quarterly)")


def _write_statement_block(ws: Worksheet, labels: list[str], df: pd.DataFrame,
                           section: str) -> None:
    if df is None or df.empty:
        return
    quarters = _quarter_columns(df, n=12)
    section_row = find_label_row(ws, section)
    if section_row is None:
        return
    # Period date headers live in the row directly under the section banner
    # (template builds this as a labeled "Period" row).
    header_row = section_row + 1
    for i, col in enumerate(quarters, start=2):
        if i > 13:
            break
        date = pd.to_datetime(col, errors="coerce")
        if pd.notna(date):
            ws.cell(row=header_row, column=i, value=date.strftime("%Y-%m-%d"))

    for label in labels:
        row = find_label_row(ws, label)
        if row is None:
            continue
        series = _row_from_yf(df, label)
        if series is None:
            continue
        for i, col in enumerate(quarters, start=2):
            if i > 13:
                break
            if col not in series.index:
                continue
            val = series[col]
            if pd.notna(val):
                ws.cell(row=row, column=i, value=float(val))


# -----------------------------------------------------------
# Section: Analysis
# -----------------------------------------------------------

def _populate_analysis(wb, is_q, bs_q, cf_q, *, asset_light: bool) -> dict:
    ws = wb["Analysis"]
    result: dict[str, Any] = {"scores": {}, "metrics": {}}

    quarters = _aligned_quarters(is_q, bs_q, cf_q, n=11)
    if not quarters:
        return result

    per_q = [_compute_ratios_for_quarter(is_q, bs_q, cf_q, q) for q in quarters]
    latest = per_q[-1]

    # Write per-quarter values across cols 2..12, latest-quarter score in col 13
    score_bands = _SCORE_BAND_MAP
    for label in _ANALYSIS_LABEL_ORDER:
        row = find_label_row(ws, label)
        if row is None:
            continue
        for i, q_values in enumerate(per_q):
            v = q_values.get(label)
            if v is not None:
                ws.cell(row=row, column=2 + i, value=v)
        if label in score_bands:
            s = scoring.score(latest.get(label), score_bands[label])
            if s is not None:
                ws.cell(row=row, column=13, value=s)

    # Industry-tag gating: blank out DSO and Inventory Turnover rows for
    # asset-light businesses.
    if asset_light:
        for label in ("DSO (Efficiency)", "Inventory Turnover (Efficiency)"):
            row = find_label_row(ws, label)
            if row is None:
                continue
            for c in range(2, 14):
                ws.cell(row=row, column=c).value = None
            ws.cell(row=row, column=13).value = "N/A"

    result["metrics"] = latest
    result["scores"] = {label: scoring.score(latest.get(label), score_bands[label])
                        for label in score_bands if latest.get(label) is not None}
    result["quarters"] = [str(q) for q in quarters]
    result["per_quarter"] = per_q
    result["fcf"] = latest.get("Free Cash Flow")
    result["total_debt"] = latest.get("Total Debt")
    result["working_capital"] = latest.get("Working Capital")
    result["sales"] = latest.get("Revenue (TTM proxy)") or 0
    result["ebit"] = latest.get("__ebit_raw")
    result["net_income"] = latest.get("__ni_raw")
    result["total_assets"] = latest.get("__total_assets_raw")
    result["total_liabilities"] = latest.get("__total_liab_raw")
    result["equity"] = latest.get("__equity_raw")
    result["cur_liab"] = latest.get("__cur_liab_raw")
    result["lt_debt"] = latest.get("__lt_debt_raw")
    result["ocf"] = latest.get("__ocf_raw")
    result["capex"] = latest.get("__capex_raw")
    return result


# Ratio label → score band lookup
_SCORE_BAND_MAP: dict[str, list] = {
    "Quality of Earnings (OCF/NI)": scoring.HIGHER_IS_BETTER,
    "EBITDA-like Cash Flow Margin": scoring.HIGHER_IS_BETTER,
    "Current Ratio (Liquidity)": scoring.CURRENT_RATIO_BANDS,
    "Quick Ratio (Liquidity)": scoring.CURRENT_RATIO_BANDS,
    "Interest Cover": scoring.INTEREST_COVER_BANDS,
    "DSO (Efficiency)": scoring.DSO_BANDS,
    "Debt-to-Equity (D/E)": scoring.LOWER_IS_BETTER,
    "Debt/Assets": scoring.LOWER_IS_BETTER,
    "Operating Margin": scoring.HIGHER_IS_BETTER,
}

# Order in which we look up rows to write per-quarter values
_ANALYSIS_LABEL_ORDER = [
    "Quality of Earnings (OCF/NI)", "EBITDA-like Cash Flow",
    "EBITDA-like Cash Flow Margin", "Proxy Δ Working Capital",
    "Free Cash Flow", "Total Debt", "Net Cash – Total Debt", "Working Capital",
    "Dependence on Debt Financing", "Negative Working Capital",
    "Massive Interest Expenses", "Accumulated Losses in RE",
    "Critical Debt Load", "Interest Burden", "Debt-to-Equity (D/E)",
    "Current Ratio (Liquidity)", "Quick Ratio (Liquidity)",
    "Inventory Turnover (Efficiency)", "DSO (Efficiency)",
    "ROIC", "ROIC – WACC", "ROCE", "CROCI", "CROCI – WACC",
    "Debt/FCF", "Debt/Net OCF", "Debt/Assets", "Interest Cover",
    "FCF/Interest", "Debt/OCF", "CAPEX % of OCF",
    "Tax Burden", "Operating Margin", "Asset Turnover", "Equity Multiplier",
    "Implied ROE",
]


def _aligned_quarters(is_q, bs_q, cf_q, n: int = 11) -> list:
    """Return the quarters (sorted oldest→newest) present in all three
    statements. Returns the last n only."""
    def _cols(df):
        if df is None or df.empty:
            return set()
        cols = list(df.columns)
        if "line" in cols:
            cols.remove("line")
        return set(cols)

    common = _cols(is_q) & _cols(bs_q) & _cols(cf_q)
    if not common:
        common = _cols(is_q) | _cols(bs_q) | _cols(cf_q)
    quarters = sorted(common, key=lambda c: pd.to_datetime(c, errors="coerce") or pd.NaT)
    return quarters[-n:]


def _q_value(df, label: str, quarter) -> float | None:
    """Lookup a single yfinance line value at a given quarter date column."""
    if df is None or df.empty:
        return None
    row = _row_from_yf(df, label)
    if row is None:
        return None
    if quarter not in row.index:
        return None
    v = row[quarter]
    if pd.isna(v):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _compute_ratios_for_quarter(is_q, bs_q, cf_q, quarter) -> dict[str, float | None]:
    """Compute the full ratio set for a single quarter."""
    rev = _q_value(is_q, "Revenue", quarter)
    cor = _q_value(is_q, "Cost of Revenue", quarter)
    ebit = _q_value(is_q, "Operating Income (EBIT)", quarter)
    ni = _q_value(is_q, "Net Income", quarter)
    int_exp = _q_value(is_q, "Interest Expense", quarter)
    pretax = _q_value(is_q, "Pretax Income", quarter)
    da = _q_value(is_q, "Depreciation & Amortization", quarter)
    ocf = _q_value(cf_q, "Operating Cash Flow", quarter)
    capex = _q_value(cf_q, "Capital Expenditures", quarter)
    tot_assets = _q_value(bs_q, "Total Assets", quarter)
    cur_assets = _q_value(bs_q, "Current Assets", quarter)
    cur_liab = _q_value(bs_q, "Current Liabilities", quarter)
    inv = _q_value(bs_q, "Inventory", quarter)
    recv = _q_value(bs_q, "Receivables", quarter)
    cash = _q_value(bs_q, "Cash & ST Investments", quarter)
    lt_debt = _q_value(bs_q, "Long-Term Debt", quarter)
    tot_liab = _q_value(bs_q, "Total Liabilities", quarter)
    equity = _q_value(bs_q, "Total Equity", quarter)

    total_debt = lt_debt
    fcf = ratios.fcf(ocf, capex)
    working_capital = None
    if cur_assets is not None and cur_liab is not None:
        working_capital = cur_assets - cur_liab
    net_cash = None
    if cash is not None or total_debt is not None:
        net_cash = (cash or 0) - (total_debt or 0)

    tb, ib, om, at, em = ratios.dupont_5step(ni, pretax, ebit, rev, tot_assets, equity)
    implied_roe = None
    if None not in (tb, ib, om, at, em):
        implied_roe = tb * ib * om * at * em

    return {
        "Quality of Earnings (OCF/NI)": ratios.quality_of_earnings(ocf, ni),
        "EBITDA-like Cash Flow": ocf,
        "EBITDA-like Cash Flow Margin": ratios.safe_div(ocf, rev),
        "Free Cash Flow": fcf,
        "Total Debt": total_debt,
        "Net Cash – Total Debt": net_cash,
        "Working Capital": working_capital,
        "Debt-to-Equity (D/E)": ratios.debt_to_equity(total_debt, equity),
        "Current Ratio (Liquidity)": ratios.current_ratio(cur_assets, cur_liab),
        "Quick Ratio (Liquidity)": ratios.quick_ratio(cur_assets, inv, cur_liab),
        "Inventory Turnover (Efficiency)": ratios.inventory_turnover(cor, inv),
        "DSO (Efficiency)": ratios.days_sales_outstanding(recv, rev, period_days=90),
        "ROIC": ratios.roic(ratios.nopat(ebit, 0.21), _safe_sum(equity, total_debt)),
        "ROCE": ratios.safe_div(ebit, _safe_sum(equity, total_debt)),
        "CROCI": ratios.croci(_safe_sum(ebit, da), _safe_sum(equity, total_debt)),
        "Debt/FCF": ratios.debt_to_fcf(total_debt, ocf, capex),
        "Debt/Net OCF": ratios.debt_to_ocf(total_debt, ocf),
        "Debt/Assets": ratios.debt_to_assets(total_debt, tot_assets),
        "Interest Cover": ratios.interest_cover(ebit, int_exp),
        "FCF/Interest": ratios.fcf_to_interest(ocf, capex, int_exp),
        "Debt/OCF": ratios.debt_to_ocf(total_debt, ocf),
        "CAPEX % of OCF": ratios.capex_pct_ocf(capex, ocf),
        "Tax Burden": tb,
        "Interest Burden": ib,
        "Operating Margin": om,
        "Asset Turnover": at,
        "Equity Multiplier": em,
        "Implied ROE": implied_roe,
        # underscore-prefixed: used by credit/valuation, not written to sheet
        "__ebit_raw": ebit,
        "__ni_raw": ni,
        "__total_assets_raw": tot_assets,
        "__total_liab_raw": tot_liab,
        "__equity_raw": equity,
        "__cur_liab_raw": cur_liab,
        "__lt_debt_raw": lt_debt,
        "__ocf_raw": ocf,
        "__capex_raw": capex,
        "Revenue (TTM proxy)": rev,
    }


# -----------------------------------------------------------
# Section: Owner Earnings & Capital Allocation
# -----------------------------------------------------------

def _populate_owner_earnings(wb, is_q, cf_q) -> None:
    ws = wb["Analysis"]
    ni_ttm = _ttm(is_q, "Net Income")
    da_ttm = _ttm(is_q, "Depreciation & Amortization")
    capex_ttm = _ttm(cf_q, "Capital Expenditures")
    rev_ttm = _ttm(is_q, "Revenue")
    ocf_ttm = _ttm(cf_q, "Operating Cash Flow")
    div_ttm = _ttm(cf_q, "Dividends Paid")
    buybacks_ttm = _ttm(cf_q, "Share Buybacks")
    debt_paydown_ttm = _ttm(cf_q, "Net Debt Issued/Repaid")

    # Maintenance CAPEX proxy: D&A (industry rule of thumb)
    maint_capex = da_ttm
    growth_capex = None
    if capex_ttm is not None and maint_capex is not None:
        growth_capex = abs(capex_ttm) - abs(maint_capex)
        growth_capex = max(growth_capex, 0)

    oe = ratios.owner_earnings(ni_ttm, da_ttm, maint_capex) if maint_capex else None
    oe_margin = ratios.safe_div(oe, rev_ttm)

    write_label_value(ws, "Owner Earnings (TTM)", oe)
    write_label_value(ws, "Owner Earnings Margin", oe_margin)
    write_label_value(ws, "Maintenance CAPEX (D&A proxy)", maint_capex)
    write_label_value(ws, "Growth CAPEX", growth_capex)

    if ocf_ttm:
        bd = ratios.capital_allocation_breakdown(
            ocf=ocf_ttm, capex=capex_ttm,
            buybacks=buybacks_ttm, dividends=div_ttm,
            debt_paydown=debt_paydown_ttm,
        )
        if bd:
            write_label_value(ws, "Capital Allocation: CAPEX %", bd["CAPEX"])
            write_label_value(ws, "Capital Allocation: Buybacks %", bd["Buybacks"])
            write_label_value(ws, "Capital Allocation: Dividends %", bd["Dividends"])
            write_label_value(ws, "Capital Allocation: M&A %", bd["M&A"])
            write_label_value(ws, "Capital Allocation: Debt Paydown %", bd["Debt Paydown"])


# -----------------------------------------------------------
# Section: Multiples Band
# -----------------------------------------------------------

def _populate_multiples_band(wb, is_q, bs_q, cf_q, px_df, market: dict) -> None:
    """Compute current vs historical bands for P/E, EV/EBITDA, EV/Sales,
    P/B, FCF Yield. Historical multiples are computed by aligning each
    fiscal quarter end with the share price at that date.
    """
    ws = wb["Analysis"]
    row = find_label_row(ws, "P/E")
    if row is None:
        return

    quarters = _aligned_quarters(is_q, bs_q, cf_q, n=44)  # up to 11 years
    if not quarters or px_df is None or px_df.empty:
        return

    px_df = px_df.copy()
    px_df["date"] = pd.to_datetime(px_df["date"])
    price_series = px_df.set_index("date")["Close"].astype(float)

    shares = market.get("shares") or 1.0

    pe_series, ev_ebitda_series, ev_sales_series = [], [], []
    pb_series, fcfy_series = [], []

    for q in quarters:
        q_ts = pd.to_datetime(q, errors="coerce")
        if pd.isna(q_ts):
            continue
        # Look for the trading day on or just before quarter-end
        if price_series.empty:
            continue
        if q_ts.tz is None and price_series.index.tz is not None:
            q_ts = q_ts.tz_localize(price_series.index.tz)
        elif q_ts.tz is not None and price_series.index.tz is None:
            q_ts = q_ts.tz_localize(None)
        prior = price_series.loc[:q_ts]
        if prior.empty:
            continue
        price_at_q = float(prior.iloc[-1])

        ttm_idx = quarters[max(0, quarters.index(q) - 3): quarters.index(q) + 1]
        ni_ttm = sum(filter(None, (_q_value(is_q, "Net Income", t) for t in ttm_idx)))
        rev_ttm = sum(filter(None, (_q_value(is_q, "Revenue", t) for t in ttm_idx)))
        ebit_ttm = sum(filter(None, (_q_value(is_q, "Operating Income (EBIT)", t) for t in ttm_idx)))
        da_ttm = sum(filter(None, (_q_value(is_q, "Depreciation & Amortization", t) for t in ttm_idx)))
        ocf_ttm = sum(filter(None, (_q_value(cf_q, "Operating Cash Flow", t) for t in ttm_idx)))
        capex_ttm = sum(filter(None, (_q_value(cf_q, "Capital Expenditures", t) for t in ttm_idx)))
        fcf_ttm = ratios.fcf(ocf_ttm, capex_ttm)

        equity_at_q = _q_value(bs_q, "Total Equity", q)
        cash_at_q = _q_value(bs_q, "Cash & ST Investments", q) or 0.0
        debt_at_q = _q_value(bs_q, "Long-Term Debt", q) or 0.0

        market_cap_at_q = price_at_q * shares
        ev_at_q = market_cap_at_q + debt_at_q - cash_at_q
        ebitda_ttm = (ebit_ttm or 0) + (da_ttm or 0)

        eps_ttm = ratios.safe_div(ni_ttm, shares)
        book_per_share = ratios.safe_div(equity_at_q, shares)

        pe_series.append(ratios.safe_div(price_at_q, eps_ttm))
        ev_ebitda_series.append(ratios.safe_div(ev_at_q, ebitda_ttm))
        ev_sales_series.append(ratios.safe_div(ev_at_q, rev_ttm))
        pb_series.append(ratios.safe_div(price_at_q, book_per_share))
        fcfy_series.append(ratios.safe_div(fcf_ttm, market_cap_at_q))

    # Build bands; current = last value in each series
    rows_data = [
        ("P/E", pe_series),
        ("EV/EBITDA", ev_ebitda_series),
        ("EV/Sales", ev_sales_series),
        ("P/B", pb_series),
        ("FCF Yield", fcfy_series),
    ]
    for label, series in rows_data:
        s = pd.Series([v for v in series if v is not None and not (isinstance(v, float) and math.isinf(v))])
        if s.empty:
            continue
        band = multiples.band_from_series(label, current=s.iloc[-1], historical=s)
        r = find_label_row(ws, label)
        if r is None:
            continue
        ws.cell(row=r, column=2).value = band.current
        ws.cell(row=r, column=3).value = band.median_3y
        ws.cell(row=r, column=4).value = band.median_5y
        ws.cell(row=r, column=5).value = band.median_10y
        ws.cell(row=r, column=6).value = band.min_5y
        ws.cell(row=r, column=7).value = band.max_5y


# -----------------------------------------------------------
# Section: Valuation
# -----------------------------------------------------------

def _populate_valuation(wb, ticker, is_q, bs_q, cf_q, px_df, market: dict,
                        force: bool = False) -> dict:
    ws = wb["Valuation"]
    result: dict[str, Any] = {}

    rev_latest = _latest_value(is_q, "Revenue")
    ebit_latest = _latest_value(is_q, "Operating Income (EBIT)")
    da_latest = _latest_value(is_q, "Depreciation & Amortization") or 0.0
    ocf_latest = _latest_value(cf_q, "Operating Cash Flow")
    capex_latest = _latest_value(cf_q, "Capital Expenditures")
    lt_debt = _latest_value(bs_q, "Long-Term Debt") or 0.0
    cash = _latest_value(bs_q, "Cash & ST Investments") or 0.0
    minority = _latest_value(bs_q, "Minority Interest") or 0.0
    shares = market.get("shares")
    price = market.get("price")

    # TTM approximations from quarterly statements (sum of last 4 quarters)
    rev_ttm = _ttm(is_q, "Revenue")
    ebit_ttm = _ttm(is_q, "Operating Income (EBIT)")
    da_ttm = _ttm(is_q, "Depreciation & Amortization")
    ocf_ttm = _ttm(cf_q, "Operating Cash Flow")
    capex_ttm = _ttm(cf_q, "Capital Expenditures")
    fcf_ttm = ratios.fcf(ocf_ttm, capex_ttm)

    # Forecast revenue from consensus when available
    rev_est = yfc.revenue_estimate(ticker, force=force)
    eps_est = yfc.earnings_estimate(ticker, force=force)

    rev_growth_y1, rev_growth_y2 = _consensus_growth(rev_est, rev_ttm)
    forecast_growths = _fade_growth(rev_growth_y1, rev_growth_y2, years=5, terminal=0.025)

    ebit_margin = ratios.safe_div(ebit_ttm, rev_ttm) or 0.10
    da_pct_rev = ratios.safe_div(da_ttm, rev_ttm) or 0.05
    capex_pct_rev = ratios.safe_div(abs(capex_ttm) if capex_ttm else None, rev_ttm) or 0.05
    nwc_pct_change_rev = 0.10

    revenues = []
    last_rev = rev_ttm or 0.0
    for g in forecast_growths:
        last_rev = last_rev * (1 + g)
        revenues.append(last_rev)

    ebits = [r * ebit_margin for r in revenues]
    da_s = [r * da_pct_rev for r in revenues]
    capexes = [r * capex_pct_rev for r in revenues]
    rev_changes = [revenues[0] - (rev_ttm or 0.0)] + [
        revenues[i] - revenues[i - 1] for i in range(1, len(revenues))
    ]
    nwc_changes = [d * nwc_pct_change_rev for d in rev_changes]

    tax_rate = market.get("tax_rate", 0.21)
    ufcf = []
    ebitda_forecast = []
    for i in range(len(revenues)):
        nopat_i = ebits[i] * (1 - tax_rate)
        ufcf_i = nopat_i + da_s[i] - capexes[i] - nwc_changes[i]
        ufcf.append(ufcf_i)
        ebitda_forecast.append(ebits[i] + da_s[i])

    # Write forecast columns Y+1..Y+5 (cols 2..6) and Terminal (col 7)
    forecast_row_labels = {
        "Revenue": revenues,
        "EBIT": ebits,
        "EBITDA": ebitda_forecast,
        "Tax": [ebits[i] * tax_rate for i in range(len(ebits))],
        "Unlevered Net Income": [ebits[i] * (1 - tax_rate) for i in range(len(ebits))],
        "+ D&A": da_s,
        "− CAPEX": [-c for c in capexes],
        "− Δ NWC": [-d for d in nwc_changes],
        "Unlevered FCF": ufcf,
    }
    for label, series in forecast_row_labels.items():
        row = find_label_row(ws, label)
        if row is None:
            continue
        for i, v in enumerate(series):
            ws.cell(row=row, column=2 + i, value=v)

    # Sensitivity grids
    rates = config.get("discount_rates", [0.03, 0.04, 0.05])
    mults = config.get("terminal_multiples", [8.0, 10.0, 12.0])
    net_debt = (lt_debt or 0.0) - (cash or 0.0)
    terminal_ebitda = ebitda_forecast[-1] if ebitda_forecast else 0.0

    grids = dcf.sensitivity_grid(
        ufcf=ufcf, terminal_ebitda=terminal_ebitda,
        discount_rates=rates, terminal_multiples=mults,
        net_debt=net_debt, minority=minority,
        shares=shares or 0.0, current_price=price or 0.0,
    )

    price_section_row = find_label_row(
        ws, "Sensitivity: Price per Share (rows = discount rate, cols = terminal multiple)"
    )
    if price_section_row:
        for r_idx, rate in enumerate(rates):
            for c_idx, mult in enumerate(mults):
                v = grids["price"].loc[rate, mult]
                if pd.notna(v):
                    ws.cell(row=price_section_row + 2 + r_idx,
                            column=2 + c_idx, value=float(v))

    upside_section_row = find_label_row(ws, "Sensitivity: Implied Upside %")
    if upside_section_row:
        for r_idx, rate in enumerate(rates):
            for c_idx, mult in enumerate(mults):
                v = grids["upside"].loc[rate, mult]
                if pd.notna(v):
                    ws.cell(row=upside_section_row + 2 + r_idx,
                            column=2 + c_idx, value=float(v))

    # Base-case point estimate uses the middle row/col of the grids
    mid_rate = rates[len(rates) // 2]
    mid_mult = mults[len(mults) // 2]
    base_price = grids["price"].loc[mid_rate, mid_mult]
    if pd.isna(base_price):
        base_price = None
    else:
        base_price = float(base_price)

    write_label_value(ws, "WACC", mid_rate)
    write_label_value(ws, "Net Debt", net_debt)
    write_label_value(ws, "Minority Interest", minority)
    write_label_value(ws, "Shares Outstanding", shares)
    write_label_value(ws, "Current Price", price)

    market_cap_now = (price or 0) * (shares or 0)
    implied_g = dcf.reverse_dcf(market_cap_now, fcf_ttm or 0.0, mid_rate)
    if implied_g is not None:
        write_label_value(ws, "Implied Growth (current price)", implied_g)
    mos = dcf.margin_of_safety(base_price, price) if (base_price and price) else None
    if mos is not None:
        write_label_value(ws, "Margin of Safety", mos)

    result["fair_value"] = base_price
    result["price"] = price
    result["upside"] = (base_price / price - 1) if (base_price and price) else None
    result["mos"] = mos
    result["implied_growth"] = implied_g
    result["net_debt"] = net_debt
    result["minority"] = minority
    result["fcf_ttm"] = fcf_ttm
    result["ufcf_forecast"] = ufcf
    result["wacc"] = mid_rate
    return result


_MAX_FORECAST_GROWTH = 0.40  # cap any single forecast year at 40% YoY


def _consensus_growth(rev_est_df, current_rev) -> tuple[float, float]:
    """Return (g_y1, g_y2) as YEAR-OVER-YEAR growth rates.

    g_y1 = (FY-current estimate / TTM) - 1
    g_y2 = (FY-next estimate / FY-current estimate) - 1   ← NOT relative to TTM

    Both are capped at _MAX_FORECAST_GROWTH to prevent compounding
    extreme analyst expectations across the full 5-year horizon. Falls
    back to 0.05 / 0.04 when consensus is missing.
    """
    if rev_est_df is None or rev_est_df.empty or current_rev in (None, 0):
        return 0.05, 0.04
    df = rev_est_df.copy()
    candidate_cols = [c for c in df.columns if "avg" in str(c).lower()]
    if not candidate_cols:
        return 0.05, 0.04
    col = candidate_cols[0]
    periods = df.iloc[:, 0].astype(str).str.lower().tolist()

    fy_current = fy_next = None
    for i, p in enumerate(periods):
        v = df.iloc[i][col]
        if pd.isna(v):
            continue
        if "+1y" in p and fy_next is None:
            fy_next = float(v)
        elif ("0y" in p or "+0y" in p) and fy_current is None:
            fy_current = float(v)

    g_y1 = (fy_current / float(current_rev) - 1) if fy_current else 0.05
    if fy_current and fy_next:
        g_y2 = fy_next / fy_current - 1
    elif fy_next:
        # Only +1y given; halve to approximate Y2-from-Y1 growth
        g_y2 = (fy_next / float(current_rev) - 1) / 2
    else:
        g_y2 = g_y1 * 0.8

    # Sanity cap. Analyst consensus for hyper-growth names regularly
    # exceeds 40% YoY; compounding that across 5 years produces fantasy
    # valuations. 40% is the upper bound of "fast-growing equity"
    # assumptions in most published DCF templates.
    g_y1 = max(min(g_y1, _MAX_FORECAST_GROWTH), -0.20)
    g_y2 = max(min(g_y2, _MAX_FORECAST_GROWTH), -0.20)
    return g_y1, g_y2


def _fade_growth(g1: float, g2: float, years: int, terminal: float) -> list[float]:
    """Linear fade from g2 down to terminal over remaining years."""
    growths = [g1, g2]
    remaining = years - 2
    if remaining <= 0:
        return growths[:years]
    step = (g2 - terminal) / (remaining + 1)
    for k in range(remaining):
        growths.append(g2 - step * (k + 1))
    return growths


def _ttm(df, label: str) -> float | None:
    row = _row_from_yf(df, label)
    if row is None:
        return None
    row = _sort_by_date_index(row)
    vals = pd.to_numeric(row, errors="coerce").dropna()
    if vals.empty:
        return None
    last4 = vals.iloc[-4:]
    return float(last4.sum())


# -----------------------------------------------------------
# Section: Credit
# -----------------------------------------------------------

def _populate_credit(wb, market: dict, is_q, bs_q, px_df) -> dict:
    ws = wb["Credit"]
    result: dict[str, Any] = {}

    equity_value = market.get("market_cap") or 0.0
    sigma_e = market.get("equity_vol") or 0.5
    rf = market.get("rf") or 0.04
    lt_debt = _latest_value(bs_q, "Long-Term Debt") or 0.0
    st_debt = _latest_value(bs_q, "Current Liabilities") or 0.0
    tot_liab = _latest_value(bs_q, "Total Liabilities") or (lt_debt + st_debt)

    write_label_value(ws, "Equity Value (E)", equity_value)
    write_label_value(ws, "Equity Volatility (σE)", sigma_e)
    write_label_value(ws, "Risk-Free Rate", rf)
    write_label_value(ws, "Long-Term Liabilities", lt_debt)
    write_label_value(ws, "Short-Term Liabilities", st_debt)
    write_label_value(ws, "Horizon (years)", 1.0)

    if equity_value > 0 and tot_liab > 0:
        k_in = kmv.KMVInputs(
            equity_value=equity_value, equity_vol=sigma_e,
            risk_free=rf, long_term_liab=lt_debt, short_term_liab=st_debt,
            horizon_years=1.0,
        )
        k = kmv.solve(k_in)
        write_label_value(ws, "Default Point (0.5·LT + ST)", k.default_point)
        write_label_value(ws, "Solved Asset Value (A)", k.asset_value)
        write_label_value(ws, "Solved Asset Volatility (σA)", k.asset_vol)
        write_label_value(ws, "Distance to Default", k.distance_to_default)
        write_label_value(ws, "EDF (Expected Default Frequency)", k.edf)
        write_label_value(ws, "Solver Squared Error", k.squared_error)
        result["edf"] = k.edf
        result["dd"] = k.distance_to_default

        m_in = merton.MertonInputs(
            equity_value=equity_value, equity_vol=sigma_e,
            total_liab=tot_liab, risk_free=rf, term_years=1.0,
        )
        m = merton.solve(m_in)
        write_label_value(ws, "Merton: Put Value", m.put_value)
        write_label_value(ws, "Merton: CDS per Year", m.cds_spread_per_year)
        write_label_value(ws, "Merton: Default Prob (Term)", m.default_prob_to_term)
        write_label_value(ws, "Merton: Annual Default Prob", m.annual_default_prob)
        result["merton_annual_pd"] = m.annual_default_prob

    # Altman
    wc = ((_latest_value(bs_q, "Current Assets") or 0) -
          (_latest_value(bs_q, "Current Liabilities") or 0))
    a_in = altman.AltmanInputs(
        working_capital=wc,
        retained_earnings=_latest_value(bs_q, "Retained Earnings") or 0,
        ebit=_latest_value(is_q, "Operating Income (EBIT)") or 0,
        market_value_equity=equity_value,
        total_liabilities=tot_liab or 1,
        sales=_ttm(is_q, "Revenue") or 0,
        total_assets=_latest_value(bs_q, "Total Assets") or 1,
    )
    a = altman.solve(a_in)
    write_label_value(ws, "X1 = WC / TA", a.x1)
    write_label_value(ws, "X2 = RE / TA", a.x2)
    write_label_value(ws, "X3 = EBIT / TA", a.x3)
    write_label_value(ws, "X4 = MVe / TL", a.x4)
    write_label_value(ws, "X5 = Sales / TA", a.x5)
    write_label_value(ws, "Z-Score", a.z_score)
    write_label_value(ws, "Bucket", a.bucket)
    result["altman_z"] = a.z_score
    result["altman_bucket"] = a.bucket

    # Hillegeist
    h_in = hillegeist.HillegeistInputs(
        working_capital=wc,
        retained_earnings=_latest_value(bs_q, "Retained Earnings") or 0,
        ebit=_latest_value(is_q, "Operating Income (EBIT)") or 0,
        market_value_equity=equity_value,
        total_liabilities=tot_liab or 1,
        total_assets=_latest_value(bs_q, "Total Assets") or 1,
    )
    h = hillegeist.solve(h_in)
    write_label_value(ws, "Hillegeist: Score", h.score)
    write_label_value(ws, "Hillegeist: Default Prob", h.default_prob)

    return result


# -----------------------------------------------------------
# Section: Cover
# -----------------------------------------------------------

def _populate_cover(wb, ticker: str, info: dict, market: dict,
                    analysis: dict, valuation: dict, credit: dict) -> None:
    ws = wb["Cover"]

    write_label_value(ws, "Ticker", ticker)
    write_label_value(ws, "Name", info.get("shortName") or info.get("longName") or "")
    write_label_value(ws, "Sector", info.get("sector") or "")
    write_label_value(ws, "Val Date", datetime.now(timezone.utc).strftime("%Y-%m-%d"))

    write_label_value(ws, "Price", market.get("price"))
    write_label_value(ws, "Fair Value (DCF)", valuation.get("fair_value"))
    write_label_value(ws, "Upside %", valuation.get("upside"))
    write_label_value(ws, "Margin of Safety", valuation.get("mos"))
    write_label_value(ws, "Reverse-DCF Growth", valuation.get("implied_growth"))

    scores = [s for s in analysis.get("scores", {}).values() if s is not None]
    if scores:
        write_label_value(ws, "Quality Score (avg)", round(sum(scores) / len(scores), 1))

    edf = credit.get("edf")
    write_label_value(ws, "EDF 1y (KMV)", edf)
    altman_text = ""
    if credit.get("altman_z") is not None:
        altman_text = f"{credit['altman_z']:.2f} ({credit.get('altman_bucket','')})"
    write_label_value(ws, "Altman Z + Bucket", altman_text)

    cal = yfc.calendar(ticker)
    earnings_raw = cal.get("Earnings Date") or cal.get("earningsDate") or ""
    earnings_str = _format_earnings_date(earnings_raw)
    write_label_value(ws, "Next Earnings", earnings_str)
    days_to = _days_to(earnings_str)
    write_label_value(ws, "Days to Earnings", days_to)

    insider = yfc.insider_transactions(ticker)
    write_label_value(ws, "Insider Net (90d)", _insider_net_90d(insider))

    # EPS revisions (3-month)
    rev = yfc.eps_revisions(ticker)
    write_label_value(ws, "3M EPS Revision", _eps_revision_3m(rev))

    # Form 4 filing count (last 90 days) — direction-agnostic but a useful
    # activity signal where yfinance insider data is sparse.
    try:
        write_label_value(ws, "Form 4 Filings (90d)",
                          edgar.form4_count(ticker, days=90))
    except Exception:
        pass

    # Top Red Flags
    red_flags = _format_red_flags(analysis.get("scores", {}))
    rf_row = find_label_row(ws, "Top Red Flags (auto-populated from Analysis tab)")
    if rf_row is not None and red_flags:
        ws.cell(row=rf_row + 1, column=1, value="\n".join(red_flags))


def _format_red_flags(scores: dict[str, int | None]) -> list[str]:
    items = [(label, s) for label, s in scores.items() if s is not None and s < 5]
    items.sort(key=lambda x: x[1])
    return [f"• {label}: {s}/10" for label, s in items[:5]]


def _format_earnings_date(raw) -> str:
    """yfinance returns earnings date as a list[date], a single date, or a
    string. After cache round-trip via str(), it may also be a literal
    bracketed-string like '[datetime.date(2026, 5, 21)]'. Normalize all
    of these to 'YYYY-MM-DD'."""
    if raw in (None, "", []):
        return ""
    if isinstance(raw, (list, tuple)) and raw:
        raw = raw[0]
    if isinstance(raw, str):
        # Strip the bracketed/datetime.date() wrapper if present.
        m = re.search(r"(\d{4})-(\d{2})-(\d{2})", raw)
        if m:
            return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
        m = re.search(r"datetime\.date\((\d+),\s*(\d+),\s*(\d+)\)", raw)
        if m:
            return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    try:
        d = pd.to_datetime(raw, errors="coerce")
        if pd.isna(d):
            return str(raw)
        return d.strftime("%Y-%m-%d")
    except Exception:
        return str(raw)


def _days_to(s: str) -> int | None:
    if not s or s == "None":
        return None
    try:
        d = pd.to_datetime(s, errors="coerce")
        if pd.isna(d):
            return None
        now = pd.Timestamp.now(tz="UTC")
        if d.tz is None:
            d = d.tz_localize("UTC")
        return int((d - now).days)
    except Exception:
        return None


def _insider_net_90d(df) -> float | None:
    if df is None or df.empty:
        return None
    if "Start Date" not in df.columns and "startDate" not in df.columns:
        return None
    date_col = "Start Date" if "Start Date" in df.columns else "startDate"
    val_col = "Value" if "Value" in df.columns else None
    if val_col is None:
        return None
    try:
        df = df.copy()
        df[date_col] = pd.to_datetime(df[date_col], errors="coerce")
        cutoff = pd.Timestamp.utcnow().tz_localize(None) - pd.Timedelta(days=90)
        recent = df[df[date_col] >= cutoff]
        if recent.empty:
            return 0.0
        return float(recent[val_col].sum())
    except Exception:
        return None


def _eps_revision_3m(df) -> float | None:
    if df is None or df.empty:
        return None
    cols = [c for c in df.columns if "3m" in str(c).lower() or "3M" in str(c)]
    if not cols:
        return None
    try:
        return float(df.iloc[0][cols[0]])
    except Exception:
        return None


# -----------------------------------------------------------
# Market_Daily populator (headless)
# -----------------------------------------------------------

_HEATMAP_SECTIONS = {
    "S&P 500 Sub-Market Performance": "sub_market",
    "S&P 500 Sub-Sector Performance (Cap-Weighted)": "sub_sector_cw",
    "S&P 500 Sub-Sector Performance (Equal-Weighted)": "sub_sector_ew",
    "Top Thematic Sectors": "thematic",
}


def populate_market_daily(workbook_path: Path, force: bool = False) -> Path:
    wb = load_workbook(workbook_path)
    ws = wb["Daily Plan"]

    # Aggregate sparkline specs across all populator subroutines, inject
    # post-save in one pass.
    sparkline_specs: dict[str, list[dict]] = {}

    _populate_positions(wb, force=force, sparkline_specs=sparkline_specs)
    _populate_summary(wb, sparkline_specs=sparkline_specs)
    _populate_earnings_calendar(wb, force=force)

    universe = config.get("universe", {})

    # SPY 5-day baseline for RS column
    spy = yfc.prices("SPY", period="6mo", force=force)
    spy_close = _close_series(spy)
    spy_5d = signals.pct_change(spy_close, 5) or 0.0

    # Sparklines for Daily Plan heatmaps. Data is written into hidden
    # cells (cols 30-59 = 30 days of closes) and a native Excel line
    # sparkline is injected post-save referencing those cells.
    HIDDEN_DATA_COL_START = 30  # AD column
    HIDDEN_DATA_COL_END = 59     # BG column
    daily_plan_sparklines: list[dict] = []

    fetched_ok = 0
    fetch_failed = []
    for section_label, universe_key in _HEATMAP_SECTIONS.items():
        anchor = find_section_anchor(ws, section_label)
        if anchor is None:
            continue
        tickers = universe.get(universe_key, [])
        for i, tkr in enumerate(tickers):
            row = anchor + 2 + i
            px = yfc.prices(tkr, period="2y", force=force)
            close = _close_series(px)
            if close is None or close.empty:
                fetch_failed.append(tkr)
                # Mark the cell so the user can see this ticker didn't pull
                ws.cell(row=row, column=16, value="(no data)").font = Font(
                    name="Calibri", size=9, italic=True, color="9CA3AF")
                continue
            fetched_ok += 1
            ws.cell(row=row, column=11, value=signals.pct_change(close, 1))
            ws.cell(row=row, column=12, value=signals.pct_change(close, 5))
            ws.cell(row=row, column=13, value=signals.pct_off_52w_high(close))
            ytd = signals.ytd_change(yfc.prices(tkr, period="1y"))
            ws.cell(row=row, column=14, value=ytd)
            five_d = signals.pct_change(close, 5)
            ws.cell(row=row, column=15,
                    value=(five_d - spy_5d) if five_d is not None else None)
            # Write Unicode-bar sparkline as cell value (always visible) AND
            # set up native line sparkline that overlays this when Excel
            # renders it. Belt + suspenders.
            last30 = close.tail(30).tolist()
            unicode_bars = block_sparkline(last30, width=20)
            spark_cell = ws.cell(row=row, column=16, value=unicode_bars)
            spark_cell.font = _SPARKLINE_FONT
            # Hidden price data for the native sparkline
            for j, v in enumerate(last30):
                ws.cell(row=row, column=HIDDEN_DATA_COL_START + j, value=float(v))
            from openpyxl.utils import get_column_letter as _col
            start_col = _col(HIDDEN_DATA_COL_START)
            end_col = _col(HIDDEN_DATA_COL_START + len(last30) - 1)
            daily_plan_sparklines.append({
                "data_range": f"'Daily Plan'!{start_col}{row}:{end_col}{row}",
                "target_cell": f"P{row}",
            })

    if fetch_failed:
        print(f"  Daily Plan: {fetched_ok} ok, {len(fetch_failed)} failed: {fetch_failed}")
    else:
        print(f"  Daily Plan: {fetched_ok} tickers ok")

    # Hide the data columns
    from openpyxl.utils import get_column_letter as _col
    for c in range(HIDDEN_DATA_COL_START, HIDDEN_DATA_COL_END + 1):
        ws.column_dimensions[_col(c)].hidden = True

    # Section 1 booleans
    if spy_close is not None and not spy_close.empty:
        anchor = find_section_anchor(ws, "1. Market Trend")
        if anchor is not None:
            ws.cell(row=anchor + 1, column=2,
                    value="YES" if signals.daily_buy_signal(spy_close) else "NO")
            ws.cell(row=anchor + 2, column=2,
                    value="YES" if signals.weekly_buy_signal(spy_close) else "NO")
            ws.cell(row=anchor + 3, column=2,
                    value="YES" if signals.above_rising_5dma(spy_close) else "NO")
            ws.cell(row=anchor + 6, column=2,
                    value=signals.market_regime(spy_close))

    # Screener: scan custom watchlist for names where both daily + weekly buy
    _populate_screener(wb, force=force)

    # Macro tab from FRED (only if API key set)
    if config.get("fred_api_key"):
        mws = wb["Macro"]
        for label, code in (("3M", "DGS3MO"), ("2Y", "DGS2"),
                            ("5Y", "DGS5"), ("10Y", "DGS10"), ("30Y", "DGS30"),
                            ("2s10s", "T10Y2Y"), ("3M10Y", "T10Y3M"),
                            ("Fed Funds Effective", "DFF"),
                            ("Unemployment", "UNRATE")):
            row = find_label_row(mws, label)
            if row is not None:
                v = fc.latest(code)
                mws.cell(row=row, column=2, value=(v / 100.0) if v is not None else None)

    if daily_plan_sparklines:
        sparkline_specs["Daily Plan"] = daily_plan_sparklines

    for sheet_name in wb.sheetnames:
        autofit_columns(wb[sheet_name])

    wb.save(workbook_path)

    # Inject native Excel line sparklines into the saved file
    if sparkline_specs:
        sparkline_injector.inject_line_sparklines(workbook_path, sparkline_specs)

    return workbook_path


def snapshot_to_archive(workbook_path: Path, archive_dir: Path | None = None) -> Path:
    """Create a timestamped copy of a workbook in the archive directory.
    Returns the snapshot path."""
    workbook_path = Path(workbook_path)
    if archive_dir is None:
        archive_dir = config.get("archive_dir")
        if archive_dir is None:
            archive_dir = workbook_path.parent / "Archive"
    archive_dir = Path(archive_dir)
    archive_dir.mkdir(parents=True, exist_ok=True)

    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M")
    stem = workbook_path.stem
    snapshot_name = f"{stem}_{ts}{workbook_path.suffix}"
    snapshot_path = archive_dir / snapshot_name
    shutil.copy(workbook_path, snapshot_path)
    return snapshot_path


def _populate_summary(wb, sparkline_specs: dict | None = None) -> None:
    """Scan the configured tickers_dir for Ticker_*.xlsx files, read the
    Cover-sheet headline tiles from each, and aggregate into the Summary
    tab on Market_Daily."""
    if "Summary" not in wb.sheetnames:
        return
    summary = wb["Summary"]

    tickers_dir = config.get("tickers_dir")
    if tickers_dir is None or not Path(tickers_dir).exists():
        tickers_dir = Path(__file__).resolve().parent.parent / "build"
    tickers_dir = Path(tickers_dir)

    files = sorted(tickers_dir.glob("Ticker_*.xlsx"))
    # Skip the template itself
    files = [f for f in files if f.name != "Ticker_TEMPLATE.xlsx"]

    # Clear existing summary rows
    for r in range(4, 54):
        for c in range(1, 14):
            summary.cell(row=r, column=c).value = None

    rows: list[dict] = []
    for f in files:
        try:
            twb = load_workbook(f, data_only=True)
        except Exception:
            continue
        if "Cover" not in twb.sheetnames:
            continue
        cover = twb["Cover"]
        label_index = _index_labels(cover)

        def _pair(label: str):
            pos = label_index.get(label)
            if pos is None:
                return None
            return cover.cell(row=pos[0], column=pos[1] + 1).value

        rows.append({
            "ticker": _pair("Ticker") or f.stem.replace("Ticker_", ""),
            "name": _pair("Name") or "",
            "sector": _pair("Sector") or "",
            "price": _pair("Price"),
            "fair_value": _pair("Fair Value (DCF)"),
            "upside": _pair("Upside %"),
            "mos": _pair("Margin of Safety"),
            "quality": _pair("Quality Score (avg)"),
            "edf": _pair("EDF 1y (KMV)"),
            "altman": _pair("Altman Z + Bucket"),
            "next_earnings": _pair("Next Earnings"),
            "eps_rev_3m": _pair("3M EPS Revision"),
            "source": f.name,
        })

    # Sort by upside descending, putting nulls last
    def _sort_key(r):
        v = r.get("upside")
        if v is None or (isinstance(v, float) and math.isnan(v)):
            return (1, 0)
        return (0, -float(v))
    rows.sort(key=_sort_key)

    for i, row in enumerate(rows[:50]):
        r = 4 + i
        summary.cell(row=r, column=1, value=row["ticker"])
        summary.cell(row=r, column=2, value=row["name"])
        summary.cell(row=r, column=3, value=row["sector"])
        summary.cell(row=r, column=4, value=_to_number(row["price"]))
        summary.cell(row=r, column=5, value=_to_number(row["fair_value"]))
        summary.cell(row=r, column=6, value=_to_number(row["upside"]))
        summary.cell(row=r, column=7, value=_to_number(row["mos"]))
        summary.cell(row=r, column=8, value=_to_number(row["quality"]))
        summary.cell(row=r, column=9, value=_to_number(row["edf"]))
        # Altman is text like "2.85 (Safe)"; preserve as-is
        summary.cell(row=r, column=10, value=row["altman"])
        summary.cell(row=r, column=11, value=row["next_earnings"])
        summary.cell(row=r, column=12, value=_to_number(row["eps_rev_3m"]))
        summary.cell(row=r, column=13, value=row["source"])
        # Belt + suspenders sparkline (Unicode + native overlay)
        try:
            px = yfc.prices(row["ticker"], period="3mo")
            close = _close_series(px)
            if close is not None and not close.empty:
                from openpyxl.utils import get_column_letter as _col
                last30 = close.tail(30).tolist()
                cell = summary.cell(row=r, column=14,
                                    value=block_sparkline(last30, width=20))
                cell.font = _SPARKLINE_FONT
                if sparkline_specs is not None:
                    HIDDEN_START = 20
                    for j, v in enumerate(last30):
                        summary.cell(row=r, column=HIDDEN_START + j,
                                     value=float(v))
                    start_col = _col(HIDDEN_START)
                    end_col = _col(HIDDEN_START + len(last30) - 1)
                    sparkline_specs.setdefault("Summary", []).append({
                        "data_range": f"Summary!{start_col}{r}:{end_col}{r}",
                        "target_cell": f"N{r}",
                    })
            else:
                summary.cell(row=r, column=14, value="(no data)").font = Font(
                    name="Calibri", size=9, italic=True, color="9CA3AF")
        except Exception:
            pass

    # Hide the data columns once
    from openpyxl.utils import get_column_letter as _col
    for c in range(20, 50):
        summary.column_dimensions[_col(c)].hidden = True


def _populate_earnings_calendar(wb, force: bool = False) -> None:
    """Pull next-earnings dates for every ticker on the Custom Watchlist
    + Positions + every Ticker_*.xlsx file. Sort by date, write to
    Earnings Calendar tab."""
    if "Earnings Calendar" not in wb.sheetnames:
        return
    ecw = wb["Earnings Calendar"]

    candidates: set[str] = set()
    if "Daily Plan" in wb.sheetnames:
        dp = wb["Daily Plan"]
        cw_anchor = find_section_anchor(
            dp, "Custom Watchlist (25 rows, user-entered)"
        )
        if cw_anchor is not None:
            for i in range(25):
                row = cw_anchor + 2 + i
                tkr = dp.cell(row=row, column=9).value
                if tkr and str(tkr).strip():
                    candidates.add(str(tkr).strip().upper())
    if "Positions" in wb.sheetnames:
        pos = wb["Positions"]
        for r in range(4, 29):
            tkr = pos.cell(row=r, column=1).value
            if tkr and str(tkr).strip():
                candidates.add(str(tkr).strip().upper())
    if "Summary" in wb.sheetnames:
        sm = wb["Summary"]
        for r in range(4, 54):
            tkr = sm.cell(row=r, column=1).value
            if tkr and str(tkr).strip():
                candidates.add(str(tkr).strip().upper())

    # Clear existing rows
    for r in range(4, 54):
        for c in range(1, 7):
            ecw.cell(row=r, column=c).value = None

    today = pd.Timestamp.now(tz="UTC").normalize()
    cutoff = today + pd.Timedelta(days=60)
    entries: list[dict] = []
    for tkr in sorted(candidates):
        cal = yfc.calendar(tkr, force=force)
        if not cal:
            continue
        raw = cal.get("Earnings Date") or cal.get("earningsDate") or ""
        date = pd.to_datetime(str(raw), errors="coerce")
        if pd.isna(date):
            continue
        if date.tz is None:
            date = date.tz_localize("UTC")
        if date < today or date > cutoff:
            continue
        entries.append({
            "date": date,
            "ticker": tkr,
            "days": int((date - today).days),
        })

    entries.sort(key=lambda e: e["date"])
    for i, e in enumerate(entries[:50]):
        r = 4 + i
        ecw.cell(row=r, column=1, value=e["date"].strftime("%Y-%m-%d"))
        ecw.cell(row=r, column=2, value=e["ticker"])
        ecw.cell(row=r, column=4, value=e["days"])
        ecw.cell(row=r, column=5, value="yfinance")


def _index_labels(ws: Worksheet, cols: tuple[int, ...] = _DEFAULT_LABEL_COLS,
                  max_row: int = 200) -> dict[str, tuple[int, int]]:
    """Build a label → (row, col) map in one pass — fast when scanning many labels."""
    out: dict[str, tuple[int, int]] = {}
    for r in range(1, max_row + 1):
        for c in cols:
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v and v not in out:
                out[v] = (r, c)
    return out


def autofit_columns(ws: Worksheet, max_width: float = 60.0,
                    min_width: float = 8.0) -> None:
    """Approximate Excel's AutoFit. Walks every cell and sets each column's
    width to the longest displayed string + small padding. Respects merged
    cells (skips them) and hidden columns (leaves untouched)."""
    from openpyxl.utils import get_column_letter

    # Skip hidden columns — those carry the sparkline price data
    hidden_cols = {
        letter for letter, dim in ws.column_dimensions.items()
        if dim.hidden
    }

    # Build the set of merged cell coordinates so we don't measure them
    merged_coords: set[tuple[int, int]] = set()
    for mr in ws.merged_cells.ranges:
        for row in range(mr.min_row, mr.max_row + 1):
            for col in range(mr.min_col, mr.max_col + 1):
                merged_coords.add((row, col))

    widths: dict[str, float] = {}
    for row_cells in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row_cells:
            if cell.value is None:
                continue
            if (cell.row, cell.column) in merged_coords:
                # Only the top-left cell of a merged range has the value;
                # skip to avoid blowing out the column width.
                continue
            letter = get_column_letter(cell.column)
            if letter in hidden_cols:
                continue
            text = _format_for_width(cell.value, cell.number_format or "")
            # 1.1 ratio approximates Calibri 11pt char width
            w = len(text) * 1.1 + 2
            if widths.get(letter, 0) < w:
                widths[letter] = w

    for letter, w in widths.items():
        w = max(min_width, min(w, max_width))
        ws.column_dimensions[letter].width = w


def _format_for_width(value, fmt: str) -> str:
    """Render value through its number_format so width calc matches display."""
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        if "%" in fmt:
            return f"{value*100:,.1f}%"
        if "$" in fmt or "USD" in fmt:
            return f"${abs(value):,.2f}"
        if "#,##0" in fmt:
            decimals = 2 if "0.00" in fmt else 0
            return f"{value:,.{decimals}f}"
        return f"{value:,.2f}"
    return str(value)


def _to_number(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and math.isnan(v):
            return None
        return v
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _populate_screener(wb, force: bool = False) -> None:
    """Scan Custom Watchlist tickers + any user-entered position tickers; flag
    those where Daily + Weekly buy signals both YES. Write to Screener sheet."""
    if "Screener" not in wb.sheetnames:
        return
    screener = wb["Screener"]
    daily_plan = wb["Daily Plan"]

    # Custom Watchlist anchor on Daily Plan
    cw_anchor = find_section_anchor(
        daily_plan, "Custom Watchlist (25 rows, user-entered)"
    )
    candidates: list[tuple[str, str]] = []
    if cw_anchor is not None:
        for i in range(25):
            row = cw_anchor + 2 + i  # +2: section + col header
            tkr = daily_plan.cell(row=row, column=9).value
            name = daily_plan.cell(row=row, column=10).value
            if tkr and str(tkr).strip():
                candidates.append((str(tkr).strip().upper(), str(name or "")))

    # Also include position tickers
    if "Positions" in wb.sheetnames:
        pos = wb["Positions"]
        for r in range(4, 29):
            tkr = pos.cell(row=r, column=1).value
            name = pos.cell(row=r, column=2).value
            if tkr and str(tkr).strip():
                candidates.append((str(tkr).strip().upper(), str(name or "")))

    # Clear existing screener rows (header is row 3; data rows from row 4)
    for r in range(4, 50):
        for c in range(1, 8):
            screener.cell(row=r, column=c).value = None

    write_row = 4
    for tkr, name in candidates:
        px = yfc.prices(tkr, period="2y", force=force)
        close = _close_series(px)
        if close is None or close.empty:
            continue
        daily = signals.daily_buy_signal(close)
        weekly = signals.weekly_buy_signal(close)
        if not (daily and weekly):
            continue
        screener.cell(row=write_row, column=1, value=tkr)
        screener.cell(row=write_row, column=2, value=name)
        screener.cell(row=write_row, column=3, value="YES")
        screener.cell(row=write_row, column=4, value="YES")
        screener.cell(row=write_row, column=5, value=signals.pct_change(close, 5))
        screener.cell(row=write_row, column=5).number_format = "0.0%;[Red]-0.0%"
        screener.cell(row=write_row, column=6, value=signals.ytd_change(px))
        screener.cell(row=write_row, column=6).number_format = "0.0%;[Red]-0.0%"
        screener.cell(row=write_row, column=7, value=signals.pct_off_52w_high(close))
        screener.cell(row=write_row, column=7).number_format = "0.0%;[Red]-0.0%"
        write_row += 1


def _populate_positions(wb, force: bool = False,
                        sparkline_specs: dict | None = None) -> None:
    """Read tickers + shares + cost basis (rows 4-28), compute current
    value, P&L, weights, and the pairwise correlation matrix."""
    if "Positions" not in wb.sheetnames:
        return
    ws = wb["Positions"]

    positions: list[dict] = []
    for r in range(4, 29):
        tkr = ws.cell(row=r, column=1).value
        if not tkr or not str(tkr).strip():
            continue
        shares = ws.cell(row=r, column=3).value
        cost = ws.cell(row=r, column=4).value
        positions.append({
            "row": r,
            "ticker": str(tkr).strip().upper(),
            "shares": float(shares or 0),
            "cost": float(cost or 0),
        })

    if not positions:
        return

    # Pull prices for each held ticker
    price_data: dict[str, pd.Series] = {}
    current_prices: dict[str, float] = {}
    for p in positions:
        px = yfc.prices(p["ticker"], period="1y", force=force)
        close = _close_series(px)
        if close is None or close.empty:
            continue
        price_data[p["ticker"]] = close
        current_prices[p["ticker"]] = float(close.iloc[-1])

    # Write per-position rows; hidden 30-day price data lives in cols 15-44
    HIDDEN_START = 15
    from openpyxl.utils import get_column_letter as _col
    total_value = 0.0
    for p in positions:
        cp = current_prices.get(p["ticker"])
        if cp is not None:
            ws.cell(row=p["row"], column=5, value=cp)
            cur_val = cp * p["shares"]
            ws.cell(row=p["row"], column=6, value=cur_val)
            total_value += cur_val
            gain = cur_val - p["shares"] * p["cost"]
            ws.cell(row=p["row"], column=7, value=gain)
            if p["shares"] * p["cost"] > 0:
                ws.cell(row=p["row"], column=8,
                        value=gain / (p["shares"] * p["cost"]))
            # Belt + suspenders sparkline: Unicode bars in the cell as a
            # fallback, native line sparkline overlays when Excel renders.
            close = price_data.get(p["ticker"])
            if close is not None:
                last30 = close.tail(30).tolist()
                spark_cell = ws.cell(row=p["row"], column=10,
                                     value=block_sparkline(last30, width=20))
                spark_cell.font = _SPARKLINE_FONT
                if sparkline_specs is not None:
                    for j, v in enumerate(last30):
                        ws.cell(row=p["row"], column=HIDDEN_START + j,
                                value=float(v))
                    start_col = _col(HIDDEN_START)
                    end_col = _col(HIDDEN_START + len(last30) - 1)
                    sparkline_specs.setdefault("Positions", []).append({
                        "data_range": f"Positions!{start_col}{p['row']}:{end_col}{p['row']}",
                        "target_cell": f"J{p['row']}",
                    })
            else:
                ws.cell(row=p["row"], column=10, value="(no data)").font = Font(
                    name="Calibri", size=9, italic=True, color="9CA3AF")

    for c in range(HIDDEN_START, HIDDEN_START + 30):
        ws.column_dimensions[_col(c)].hidden = True

    # Write weights and total
    if total_value > 0:
        for p in positions:
            cur_val = ws.cell(row=p["row"], column=6).value or 0
            ws.cell(row=p["row"], column=9, value=cur_val / total_value)
        ws.cell(row=30, column=6, value=total_value)

    # Pairwise correlation matrix
    held = [p["ticker"] for p in positions if p["ticker"] in price_data]
    if len(held) < 2:
        return
    returns = pd.DataFrame({
        t: price_data[t].pct_change()
        for t in held
    }).dropna(how="all").tail(90)
    corr = returns.corr()

    corr_anchor = 33
    # Column headers
    for i, t in enumerate(held):
        ws.cell(row=corr_anchor + 1, column=2 + i, value=t).font = _BOLD_FONT
        ws.cell(row=corr_anchor + 2 + i, column=1, value=t).font = _BOLD_FONT
    # Matrix values
    for i, t_row in enumerate(held):
        for j, t_col in enumerate(held):
            v = corr.iloc[i, j] if t_row in corr.index and t_col in corr.columns else None
            if v is not None and not (isinstance(v, float) and math.isnan(v)):
                ws.cell(row=corr_anchor + 2 + i, column=2 + j, value=float(v))
                ws.cell(row=corr_anchor + 2 + i, column=2 + j).number_format = "0.00"


def _close_series(px_df) -> pd.Series | None:
    if px_df is None or px_df.empty:
        return None
    df = px_df.copy()
    df["date"] = pd.to_datetime(df["date"])
    return df.set_index("date")["Close"].astype(float)


# -----------------------------------------------------------
# Misc helpers
# -----------------------------------------------------------

def _safe_float(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _safe_sum(*vals) -> float | None:
    parts = [v for v in vals if v is not None]
    if not parts:
        return None
    return sum(parts)
