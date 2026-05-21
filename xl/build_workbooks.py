"""Build the Market_Daily.xlsx and Ticker_TEMPLATE.xlsx templates."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.styles import Font, PatternFill

import config
from xl import styles as S


def _set_header(ws, text: str, span_cols: int = 16):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span_cols)
    c = ws.cell(row=1, column=1, value=text)
    c.fill = S.HEADER_FILL
    c.font = S.HEADER_FONT
    c.alignment = S.CENTER
    ws.row_dimensions[1].height = 26


def _section(ws, row: int, col: int, text: str, span: int = 1):
    if span > 1:
        ws.merge_cells(start_row=row, start_column=col, end_row=row,
                       end_column=col + span - 1)
    c = ws.cell(row=row, column=col, value=text)
    c.fill = S.SECTION_FILL
    c.font = S.SECTION_FONT
    c.alignment = S.LEFT
    return row + 1


def _label_value(ws, row: int, col: int, label: str, value=None,
                 fmt: str | None = None, input_cell: bool = False):
    a = ws.cell(row=row, column=col, value=label)
    a.font = S.LABEL_FONT
    a.alignment = S.LEFT
    b = ws.cell(row=row, column=col + 1, value=value)
    b.font = S.BODY_FONT
    b.alignment = S.RIGHT
    if fmt:
        b.number_format = fmt
    if input_cell:
        b.fill = S.INPUT_FILL
    return row + 1


def _col_widths(ws, widths: dict[int, float]):
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w


# -------------------------------------------------------------------
# Market_Daily.xlsx
# -------------------------------------------------------------------

HEATMAP_COLS = ["Ticker", "Name", "% Daily", "% 5D", "% Off 52w-Hi", "% YTD", "RS-SPY", "Trend 30d"]
SUB_MARKET = [
    ("IVE", "S&P 500 Large Cap Value"),
    ("IJS", "S&P 600 Small Cap Value"),
    ("IJJ", "S&P 400 Mid Cap Value"),
    ("IJT", "S&P 600 Small Cap Growth"),
    ("IVW", "S&P 500 Large Cap Growth"),
    ("IJK", "S&P 400 Mid Cap Growth"),
]
SUB_SECTOR_CW = [
    ("XLE", "Energy"), ("XLP", "Consumer Staples"), ("XLV", "Healthcare"),
    ("XLF", "Financials"), ("XLRE", "Real Estate"), ("XLU", "Utilities"),
    ("XLC", "Communication Services"), ("SPY", "S&P 500 ETF"),
    ("XLK", "Technology"), ("XLI", "Industrials"),
    ("XLY", "Consumer Discretionary"), ("XLB", "Materials"),
]
SUB_SECTOR_EW = [
    ("RSPG", "Energy EW"), ("RSPS", "Staples EW"), ("RSPF", "Financials EW"),
    ("RSPH", "Healthcare EW"), ("RSPU", "Utilities EW"), ("RSPR", "Real Estate EW"),
    ("RSPC", "Comm Svcs EW"), ("RSP", "S&P 500 EW"),
    ("RSPD", "Consumer Disc EW"), ("RSPT", "Technology EW"),
    ("RSPN", "Industrials EW"), ("RSPM", "Materials EW"),
]
THEMATIC = [
    ("UNG", "US Nat Gas"), ("USO", "US Oil"), ("XOP", "Oil & Gas Exp"),
    ("IXC", "Global Energy"), ("IGV", "Expanded Tch-Sftwr"),
    ("IHI", "US Med Dev"), ("CLOU", "Global Cloud Computing"),
    ("XSW", "S&P Sftware & Svc"), ("KIE", "S&P Insurance"),
    ("OIH", "Oil Services"),
]


def _heatmap_block(ws, start_row: int, title: str, rows: list[tuple[str, str]]) -> int:
    r = _section(ws, start_row, 9, title, span=len(HEATMAP_COLS))
    for i, col in enumerate(HEATMAP_COLS, start=9):
        c = ws.cell(row=r, column=i, value=col)
        c.fill = S.SUBSECTION_FILL
        c.font = S.LABEL_FONT
        c.alignment = S.CENTER
        c.border = S.BOX
    r += 1

    for tkr, name in rows:
        ws.cell(row=r, column=9, value=tkr).font = S.LABEL_FONT
        ws.cell(row=r, column=10, value=name).font = S.BODY_FONT
        for i in range(11, 16):
            cell = ws.cell(row=r, column=i)
            cell.number_format = S.PCT_FMT
            cell.alignment = S.RIGHT
            cell.border = S.BOX
        # Trend sparkline cell (filled by populator)
        sp_cell = ws.cell(row=r, column=16)
        sp_cell.font = Font(name="Consolas", size=11)
        sp_cell.alignment = S.LEFT
        sp_cell.border = S.BOX
        r += 1

    # heatmap color scale on the % columns only (K..O), not the sparkline (P)
    rng = f"K{r - len(rows)}:O{r - 1}"
    rule = ColorScaleRule(
        start_type="num", start_value=-0.05, start_color="EF4444",
        mid_type="num", mid_value=0, mid_color="FFFFFF",
        end_type="num", end_value=0.05, end_color="10B981",
    )
    ws.conditional_formatting.add(rng, rule)
    return r + 1


def build_market_daily(out_path: Path) -> Path:
    wb = Workbook()

    # --- Daily Plan ---
    ws = wb.active
    ws.title = "Daily Plan"
    _set_header(ws, 'Daily Market Plan – "If you fail to plan, you are planning to fail" – Benjamin Franklin', 16)
    _col_widths(ws, {1: 28, 2: 14, 3: 14, 4: 14, 5: 14, 6: 14, 7: 14, 8: 2,
                     9: 8, 10: 32, 11: 10, 12: 10, 13: 12, 14: 10, 15: 10, 16: 24})

    # Left column blocks
    r = 3
    r = _section(ws, r, 1, "1. Market Trend", span=7)
    for label in ("Daily Buy Signal? (price > 10/20 DMA, 10>20)",
                  "Weekly Buy Signal? (price > 10/20 WMA, 10>20)",
                  "Price > Rising 5-DMA?",
                  "Recent trade traction?",
                  "Net 52w High/Low trending positive?"):
        r = _label_value(ws, r, 1, label, value="", input_cell=False)
    r = _label_value(ws, r, 1, "Market Regime:", value="", input_cell=False)
    r += 1

    r = _section(ws, r, 1, "2. Primary Watchlist & Focus Themes", span=7)
    ws.cell(row=r, column=1, value="Focus Themes (free text):").font = S.MUTED_FONT
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 4, end_column=7)
    ws.cell(row=r, column=1).alignment = S.TOP_LEFT
    ws.cell(row=r, column=1).fill = S.INPUT_FILL
    r += 6

    r = _section(ws, r, 1, "3. Situational Awareness", span=7)
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 6, end_column=7)
    ws.cell(row=r, column=1).alignment = S.TOP_LEFT
    ws.cell(row=r, column=1).fill = S.INPUT_FILL
    r += 8

    r = _section(ws, r, 1, "4. Emotional Analysis", span=7)
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 8, end_column=7)
    ws.cell(row=r, column=1).alignment = S.TOP_LEFT
    ws.cell(row=r, column=1).fill = S.INPUT_FILL

    # Right column blocks
    r2 = 3
    r2 = _section(ws, r2, 9, "Market Commentary", span=len(HEATMAP_COLS))
    ws.merge_cells(start_row=r2, start_column=9, end_row=r2 + 4, end_column=15)
    ws.cell(row=r2, column=9).alignment = S.TOP_LEFT
    ws.cell(row=r2, column=9).fill = S.INPUT_FILL
    r2 += 6

    r2 = _section(ws, r2, 9, "Macro Risk Strip (VIX, MOVE, HYG-LQD, DXY, 2s10s)", span=len(HEATMAP_COLS))
    for i, label in enumerate(["VIX", "MOVE", "HYG-LQD", "DXY", "2s10s"]):
        ws.cell(row=r2, column=9 + i, value=label).font = S.LABEL_FONT
        ws.cell(row=r2, column=9 + i).alignment = S.CENTER
    r2 += 1
    for i in range(5):
        ws.cell(row=r2, column=9 + i).number_format = S.NUM_FMT
        ws.cell(row=r2, column=9 + i).alignment = S.CENTER
    r2 += 2

    r2 = _heatmap_block(ws, r2, "S&P 500 Sub-Market Performance", SUB_MARKET)
    r2 = _heatmap_block(ws, r2, "S&P 500 Sub-Sector Performance (Cap-Weighted)", SUB_SECTOR_CW)
    r2 = _heatmap_block(ws, r2, "S&P 500 Sub-Sector Performance (Equal-Weighted)", SUB_SECTOR_EW)
    r2 = _heatmap_block(ws, r2, "Top Thematic Sectors", THEMATIC)
    r2 = _heatmap_block(ws, r2, "Custom Watchlist (25 rows, user-entered)",
                       [("", "") for _ in range(25)])

    ws.freeze_panes = "I3"

    # --- Breadth tab ---
    bws = wb.create_sheet("Breadth")
    _set_header(bws, "Watchlist Breadth & Internals", 10)
    _col_widths(bws, {1: 34, 2: 14, 3: 40})
    rb = 3
    rb = _section(bws, rb, 1, "Breadth — computed across your watchlist universe", span=3)
    for label, fmt in (("Tickers with data", S.INT_FMT),
                       ("% above 50-DMA", S.PCT_FMT),
                       ("% above 200-DMA", S.PCT_FMT),
                       ("New 52-week Highs", S.INT_FMT),
                       ("New 52-week Lows", S.INT_FMT),
                       ("Advancers (today)", S.INT_FMT),
                       ("Decliners (today)", S.INT_FMT),
                       ("Advance/Decline ratio", S.NUM_FMT),
                       ("% with Daily Buy Signal", S.PCT_FMT),
                       ("% with Weekly Buy Signal", S.PCT_FMT)):
        bws.cell(row=rb, column=1, value=label).font = S.LABEL_FONT
        bws.cell(row=rb, column=2).number_format = fmt
        bws.cell(row=rb, column=2).alignment = S.RIGHT
        rb += 1

    # --- Watchlist tab (full list, one row per ticker) ---
    wlw = wb.create_sheet("Watchlist")
    _set_header(wlw, "Watchlist — full list (edit watchlist.txt to change)", 10)
    _col_widths(wlw, {1: 10, 2: 30, 3: 11, 4: 11, 5: 13, 6: 11, 7: 11,
                      8: 11, 9: 11, 10: 26})
    wl_headers = ["Ticker", "Name", "% Daily", "% 5D", "% Off 52w-Hi",
                  "% YTD", "RS-SPY", "Daily Buy", "Weekly Buy", "Trend 30d"]
    for i, h in enumerate(wl_headers):
        c = wlw.cell(row=3, column=i + 1, value=h)
        c.fill = S.SUBSECTION_FILL
        c.font = S.LABEL_FONT
        c.alignment = S.CENTER
        c.border = S.BOX
    # 300 data rows — plenty for a few hundred tickers
    for r in range(4, 304):
        for c in range(3, 8):
            wlw.cell(row=r, column=c).number_format = S.PCT_FMT
        wlw.cell(row=r, column=10).font = Font(name="Consolas", size=11)
    # Color scale on % columns
    rule = ColorScaleRule(
        start_type="num", start_value=-0.05, start_color="EF4444",
        mid_type="num", mid_value=0, mid_color="FFFFFF",
        end_type="num", end_value=0.05, end_color="10B981",
    )
    wlw.conditional_formatting.add("C4:G303", rule)
    # AutoFilter so the user can sort by any column
    wlw.auto_filter.ref = "A3:J303"
    wlw.freeze_panes = "A4"

    # --- Macro tab ---
    mws = wb.create_sheet("Macro")
    _set_header(mws, "Macro Reference", 8)
    _col_widths(mws, {1: 32, 2: 14, 3: 14, 4: 14, 5: 30})
    rm = 3

    def _macro_section(start_row, title, rows):
        """Section banner + Current/Previous/Change header + labelled rows."""
        r = _section(mws, start_row, 1, title, span=4)
        for i, h in enumerate(("", "Current", "Previous", "Change")):
            c = mws.cell(row=r, column=1 + i, value=h)
            if h:
                c.fill = S.SUBSECTION_FILL
                c.font = S.LABEL_FONT
                c.alignment = S.CENTER
        r += 1
        for label, fmt in rows:
            mws.cell(row=r, column=1, value=label).font = S.LABEL_FONT
            for c in (2, 3, 4):
                mws.cell(row=r, column=c).number_format = fmt
                mws.cell(row=r, column=c).alignment = S.RIGHT
            r += 1
        return r + 1

    rm = _macro_section(rm, "US Treasury Curve (Current vs 1 week ago)", [
        ("3M", S.PCT_FMT), ("2Y", S.PCT_FMT), ("5Y", S.PCT_FMT),
        ("10Y", S.PCT_FMT), ("30Y", S.PCT_FMT)])
    rm = _macro_section(rm, "Curve Spreads (Current vs 1 week ago)", [
        ("2s10s", S.PCT_FMT), ("3M10Y", S.PCT_FMT)])
    rm = _macro_section(rm, "Macro Indicators", [
        ("Fed Funds Effective", S.PCT_FMT),
        ("CPI YoY", S.PCT_FMT),
        ("PPI YoY", S.PCT_FMT),
        ("Unemployment", S.PCT_FMT),
        ("Nonfarm Payrolls (chg, 000s)", S.NUM_FMT),
        ("GDP Nowcast (Atlanta Fed)", S.PCT_FMT)])
    rm = _section(mws, rm, 1, "Calendar", span=4)
    mws.cell(row=rm, column=1, value="Next FOMC date").font = S.LABEL_FONT
    rm += 1
    mws.cell(row=rm, column=1, value="Days to FOMC").font = S.LABEL_FONT
    mws.cell(row=rm, column=2).number_format = S.INT_FMT

    # --- Screener tab ---
    sws = wb.create_sheet("Screener")
    _set_header(sws, "Auto-Screener: Custom Watchlist names where Daily AND Weekly Buy = YES", 12)
    _col_widths(sws, {1: 10, 2: 32, 3: 12, 4: 12, 5: 12, 6: 12, 7: 12})
    for i, h in enumerate(["Ticker", "Name", "Daily Buy", "Weekly Buy",
                            "% 5D", "% YTD", "Off 52w-Hi"]):
        c = sws.cell(row=3, column=i + 1, value=h)
        c.fill = S.SUBSECTION_FILL
        c.font = S.LABEL_FONT
        c.alignment = S.CENTER
        c.border = S.BOX

    # --- Watchlist Summary (aggregates Cover tiles from every Ticker_*.xlsx) ---
    sws = wb.create_sheet("Summary")
    _set_header(sws, "Watchlist Summary — aggregates Cover sheet from every Ticker workbook", 14)
    _col_widths(sws, {1: 10, 2: 24, 3: 16, 4: 12, 5: 12, 6: 10, 7: 10, 8: 10, 9: 10, 10: 12, 11: 14, 12: 10, 13: 16, 14: 24})
    headers = ["Ticker", "Name", "Sector", "Price", "Fair Value",
               "Upside %", "MoS %", "Quality", "EDF", "Altman Z",
               "Next Earnings", "3M EPS Rev", "Source File", "Trend 30d"]
    for i, h in enumerate(headers):
        c = sws.cell(row=3, column=i + 1, value=h)
        c.fill = S.SUBSECTION_FILL
        c.font = S.LABEL_FONT
        c.alignment = S.CENTER
        c.border = S.BOX
    for r in range(4, 54):
        for c in range(1, 15):
            sws.cell(row=r, column=c).border = S.BOX
        sws.cell(row=r, column=4).number_format = S.USD_FMT
        sws.cell(row=r, column=5).number_format = S.USD_FMT
        sws.cell(row=r, column=6).number_format = S.PCT_FMT
        sws.cell(row=r, column=7).number_format = S.PCT_FMT
        sws.cell(row=r, column=8).number_format = S.NUM_FMT
        sws.cell(row=r, column=9).number_format = S.PCT_FMT
        sws.cell(row=r, column=10).number_format = S.NUM_FMT
        sws.cell(row=r, column=12).number_format = S.PCT_FMT
        sws.cell(row=r, column=14).font = Font(name="Consolas", size=11)

    # Conditional formatting on upside col (F), MoS col (G), and Quality (H)
    rule_pos = ColorScaleRule(
        start_type="num", start_value=-0.5, start_color="EF4444",
        mid_type="num", mid_value=0, mid_color="FFFFFF",
        end_type="num", end_value=0.5, end_color="10B981",
    )
    sws.conditional_formatting.add("F4:G53", rule_pos)
    rule_quality = ColorScaleRule(
        start_type="num", start_value=0, start_color="EF4444",
        mid_type="num", mid_value=5, mid_color="FBBF24",
        end_type="num", end_value=10, end_color="10B981",
    )
    sws.conditional_formatting.add("H4:H53", rule_quality)
    rule_edf = ColorScaleRule(
        start_type="num", start_value=0, start_color="10B981",
        mid_type="num", mid_value=0.05, mid_color="FBBF24",
        end_type="num", end_value=0.2, end_color="EF4444",
    )
    sws.conditional_formatting.add("I4:I53", rule_edf)

    # --- Earnings Calendar (next 60 days, watchlist + positions) ---
    ecw = wb.create_sheet("Earnings Calendar")
    _set_header(ecw, "Earnings Calendar — next 60 days, sorted by date", 6)
    _col_widths(ecw, {1: 12, 2: 10, 3: 24, 4: 14, 5: 12, 6: 14})
    headers = ["Date", "Ticker", "Name", "Days From Now", "Source", "Note"]
    for i, h in enumerate(headers):
        c = ecw.cell(row=3, column=i + 1, value=h)
        c.fill = S.SUBSECTION_FILL
        c.font = S.LABEL_FONT
        c.alignment = S.CENTER
        c.border = S.BOX
    for r in range(4, 54):
        for c in range(1, 7):
            ecw.cell(row=r, column=c).border = S.BOX
        ecw.cell(row=r, column=4).number_format = S.INT_FMT

    # --- Lessons Learned ---
    lws = wb.create_sheet("Lessons")
    _set_header(lws, "Lessons Learned (append-only)", 8)
    _col_widths(lws, {1: 12, 2: 10, 3: 14, 4: 40, 5: 40, 6: 40})
    headers = ["Date", "Ticker", "Action", "Reasoning", "Outcome", "Lesson"]
    for i, h in enumerate(headers):
        c = lws.cell(row=3, column=i + 1, value=h)
        c.fill = S.SUBSECTION_FILL
        c.font = S.LABEL_FONT
        c.alignment = S.CENTER
        c.border = S.BOX
    # Reserve 100 rows for entries with light formatting
    for r in range(4, 104):
        for c in range(1, 7):
            lws.cell(row=r, column=c).border = S.BOX
            lws.cell(row=r, column=c).alignment = S.TOP_LEFT

    # --- Positions ---
    pws = wb.create_sheet("Positions")
    _set_header(pws, "Portfolio Positions & Correlations", 12)
    _col_widths(pws, {1: 10, 2: 24, 3: 12, 4: 12, 5: 12, 6: 14, 7: 14, 8: 12, 9: 12})

    headers = ["Ticker", "Name", "Shares", "Cost Basis", "Current Price",
               "Current Value", "Gain/Loss $", "Gain/Loss %", "Weight %",
               "Trend 30d"]
    for i, h in enumerate(headers):
        c = pws.cell(row=3, column=i + 1, value=h)
        c.fill = S.SUBSECTION_FILL
        c.font = S.LABEL_FONT
        c.alignment = S.CENTER
        c.border = S.BOX
    pws.column_dimensions[get_column_letter(10)].width = 24
    # 25 rows for positions
    for r in range(4, 29):
        for c in (1, 2, 3, 4):
            pws.cell(row=r, column=c).fill = S.INPUT_FILL
        pws.cell(row=r, column=3).number_format = S.INT_FMT
        pws.cell(row=r, column=4).number_format = S.USD_FMT
        pws.cell(row=r, column=5).number_format = S.USD_FMT
        pws.cell(row=r, column=6).number_format = S.USD_FMT
        pws.cell(row=r, column=7).number_format = S.USD_FMT
        pws.cell(row=r, column=8).number_format = S.PCT_FMT
        pws.cell(row=r, column=9).number_format = S.PCT_FMT
        for c in range(1, 11):
            pws.cell(row=r, column=c).border = S.BOX
        pws.cell(row=r, column=10).font = Font(name="Consolas", size=11)

    # Totals row
    tot_row = 30
    pws.cell(row=tot_row, column=2, value="Total").font = S.LABEL_FONT
    pws.cell(row=tot_row, column=6).number_format = S.USD_FMT
    pws.cell(row=tot_row, column=7).number_format = S.USD_FMT

    # Pairwise correlation matrix starts at row 33
    corr_anchor = 33
    pws.cell(row=corr_anchor, column=1, value="Pairwise Correlation (90d daily returns)").font = S.SECTION_FONT
    pws.cell(row=corr_anchor, column=1).fill = S.SECTION_FILL
    pws.merge_cells(start_row=corr_anchor, start_column=1,
                    end_row=corr_anchor, end_column=12)
    # 25x25 area for the matrix, gets filled by populator

    rule = ColorScaleRule(
        start_type="num", start_value=-1.0, start_color="2563EB",  # blue
        mid_type="num", mid_value=0.0, mid_color="FFFFFF",
        end_type="num", end_value=1.0, end_color="EF4444",         # red
    )
    pws.conditional_formatting.add(f"B{corr_anchor + 2}:Z{corr_anchor + 27}", rule)

    # --- Settings ---
    cws = wb.create_sheet("Settings")
    _set_header(cws, "Settings (managed via settings.local.json)", 4)
    _col_widths(cws, {1: 32, 2: 60})
    _label_value(cws, 3, 1, "FRED API Key", value="(loaded from settings.local.json)")
    _label_value(cws, 4, 1, "SEC User-Agent", value="(loaded from settings.local.json)")
    _label_value(cws, 5, 1, "Drive root", value="(loaded from settings.local.json)")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


# -------------------------------------------------------------------
# Ticker_TEMPLATE.xlsx
# -------------------------------------------------------------------

def build_ticker_template(out_path: Path) -> Path:
    wb = Workbook()

    # --- Cover ---
    ws = wb.active
    ws.title = "Cover"
    _set_header(ws, "Ticker Analysis Dashboard", 12)
    _col_widths(ws, {1: 22, 2: 18, 3: 4, 4: 22, 5: 18, 6: 4, 7: 22, 8: 18, 9: 4, 10: 22, 11: 18})

    r = 3
    _label_value(ws, r, 1, "Ticker", value="", input_cell=True)
    _label_value(ws, r, 4, "Name", value="")
    _label_value(ws, r, 7, "Sector", value="")
    _label_value(ws, r, 10, "Val Date", value="")
    r += 2

    r = _section(ws, r, 1, "Headline Metrics", span=12)
    _label_value(ws, r, 1, "Price", fmt=S.USD_FMT)
    _label_value(ws, r, 4, "Fair Value (DCF)", fmt=S.USD_FMT)
    _label_value(ws, r, 7, "Upside %", fmt=S.PCT_FMT)
    _label_value(ws, r, 10, "Margin of Safety", fmt=S.PCT_FMT)
    r += 1
    _label_value(ws, r, 1, "Reverse-DCF Growth", fmt=S.PCT_FMT)
    _label_value(ws, r, 4, "Quality Score (avg)", fmt=S.SCORE_FMT)
    _label_value(ws, r, 7, "EDF 1y (KMV)", fmt=S.PCT_FMT)
    _label_value(ws, r, 10, "Altman Z + Bucket", value="")
    r += 1
    _label_value(ws, r, 1, "Next Earnings", value="")
    _label_value(ws, r, 4, "Days to Earnings", fmt=S.INT_FMT)
    _label_value(ws, r, 7, "Insider Net (90d)", fmt=S.USD_FMT)
    _label_value(ws, r, 10, "Short Int % Float", fmt=S.PCT_FMT)
    r += 1
    _label_value(ws, r, 1, "3M EPS Revision", fmt=S.PCT_FMT)
    _label_value(ws, r, 4, "Form 4 Filings (90d)", fmt=S.INT_FMT)
    r += 2

    r = _section(ws, r, 1, "Top Red Flags (auto-populated from Analysis tab)", span=12)
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 3, end_column=12)
    ws.cell(row=r, column=1).alignment = S.TOP_LEFT
    r += 5

    r = _section(ws, r, 1, "Theses", span=12)
    for label in ("Long thesis", "Risks", "Pre-mortem (I'd know I'm wrong if...)"):
        ws.cell(row=r, column=1, value=label).font = S.LABEL_FONT
        ws.merge_cells(start_row=r, start_column=2, end_row=r + 2, end_column=12)
        ws.cell(row=r, column=2).alignment = S.TOP_LEFT
        ws.cell(row=r, column=2).fill = S.INPUT_FILL
        r += 3

    # --- Market ---
    mws = wb.create_sheet("Market")
    _set_header(mws, "Market & Macro Inputs", 8)
    _col_widths(mws, {1: 28, 2: 14, 3: 6, 4: 28, 5: 14})
    rm = 3
    rm = _section(mws, rm, 1, "Live Quote", span=2)
    for label, fmt in (("Price", S.USD_FMT), ("Shares Outstanding", S.INT_FMT),
                       ("Market Cap", S.USD_FMT), ("Dividend Yield", S.PCT_FMT),
                       ("52w High", S.USD_FMT), ("52w Low", S.USD_FMT)):
        rm = _label_value(mws, rm, 1, label, fmt=fmt)
    rm += 1
    rm = _section(mws, rm, 1, "Volatility", span=2)
    for label in ("Historical Vol (30d)", "Historical Vol (90d)",
                  "Historical Vol (365d)", "ATM Implied Vol"):
        rm = _label_value(mws, rm, 1, label, fmt=S.PCT_FMT)

    rm2 = 3
    rm2 = _section(mws, rm2, 4, "CAPM Block", span=2)
    for label, fmt, inp in (("Risk-Free Rate (10Y)", S.PCT_FMT, False),
                            ("Market Risk Premium", S.PCT_FMT, True),
                            ("Country Risk Premium", S.PCT_FMT, True),
                            ("Beta", S.NUM_FMT, False),
                            ("Cost of Equity (Re)", S.PCT_FMT, False),
                            ("Cost of Debt (Rd)", S.PCT_FMT, False),
                            ("Tax Rate", S.PCT_FMT, True),
                            ("WACC", S.PCT_FMT, False)):
        rm2 = _label_value(mws, rm2, 4, label, fmt=fmt, input_cell=inp)

    rm2 += 1
    rm2 = _section(mws, rm2, 4, "Debt Market", span=2)
    for label, fmt, inp in (("Highest debt price (% par)", S.PCT_FMT, True),
                            ("Bond Yield (proxy Rd)", S.PCT_FMT, True),
                            ("CDS Quote (if available)", S.PCT_FMT, True)):
        rm2 = _label_value(mws, rm2, 4, label, fmt=fmt, input_cell=inp)

    # --- Fin Stat ---
    fws = wb.create_sheet("Fin Stat")
    _set_header(fws, "Financial Statements (quarterly + annual)", 14)
    _col_widths(fws, {1: 40})
    for c in range(2, 14):
        fws.column_dimensions[get_column_letter(c)].width = 14
    def _stmt_block(start_row, section_label, labels):
        r = _section(fws, start_row, 1, section_label, span=14)
        # Period-date header row (filled by populator at refresh time)
        fws.cell(row=r, column=1, value="Period").font = S.MUTED_FONT
        for c in range(2, 14):
            cell = fws.cell(row=r, column=c)
            cell.fill = S.SUBSECTION_FILL
            cell.font = S.LABEL_FONT
            cell.alignment = S.CENTER
        r += 1
        for label in labels:
            fws.cell(row=r, column=1, value=label).font = S.LABEL_FONT
            for c in range(2, 14):
                fws.cell(row=r, column=c).number_format = S.NUM_FMT
            r += 1
        return r + 1

    rf = 3
    is_labels = ["Revenue", "Cost of Revenue", "Gross Profit", "SG&A", "R&D",
                 "Depreciation & Amortization", "Other Opex",
                 "Operating Income (EBIT)", "Interest Expense",
                 "Other Non-Op Income", "Pretax Income",
                 "Provision for Taxes", "Net Income"]
    is_section_row = rf
    is_period_row = rf + 1
    is_first_label_row = rf + 2
    rf = _stmt_block(rf, "Income Statement (Quarterly)", is_labels)
    is_row_of = {label: is_first_label_row + i for i, label in enumerate(is_labels)}
    rf = _stmt_block(rf, "Balance Sheet (Quarterly)", [
        "Cash & ST Investments", "Receivables", "Inventory",
        "Current Assets", "Total Assets",
        "Current Liabilities", "Long-Term Debt", "Total Liabilities",
        "Retained Earnings", "Total Equity", "Minority Interest"])
    rf = _stmt_block(rf, "Cash Flow (Quarterly)", [
        "Operating Cash Flow", "Capital Expenditures",
        "Free Cash Flow", "Dividends Paid", "Share Buybacks",
        "Net Debt Issued/Repaid"])

    # --- Analysis ---
    ws_a = wb.create_sheet("Analysis")
    _set_header(ws_a, "Company Analysis (ratios + 0-10 scoring)", 14)
    _col_widths(ws_a, {1: 36})
    for c in range(2, 14):
        ws_a.column_dimensions[get_column_letter(c)].width = 12
    ra = 3
    ra = _section(ws_a, ra, 1, "CFO / Earnings Quality", span=14)
    for label in ("Quality of Earnings (OCF/NI)", "EBITDA-like Cash Flow",
                  "EBITDA-like Cash Flow Margin", "Proxy Δ Working Capital",
                  "Free Cash Flow", "Total Debt", "Net Cash – Total Debt",
                  "Working Capital"):
        ws_a.cell(row=ra, column=1, value=label).font = S.LABEL_FONT
        for c in range(2, 13):
            ws_a.cell(row=ra, column=c).number_format = S.NUM_FMT
        ws_a.cell(row=ra, column=13, value="").number_format = S.SCORE_FMT
        ws_a.cell(row=ra, column=13).alignment = S.CENTER
        ra += 1

    ra += 1
    ra = _section(ws_a, ra, 1, "Financial Risk", span=14)
    for label in ("Dependence on Debt Financing", "Negative Working Capital",
                  "Massive Interest Expenses", "Accumulated Losses in RE",
                  "Critical Debt Load", "Interest Burden",
                  "Debt-to-Equity (D/E)",
                  "Current Ratio (Liquidity)", "Quick Ratio (Liquidity)",
                  "Inventory Turnover (Efficiency)", "DSO (Efficiency)"):
        ws_a.cell(row=ra, column=1, value=label).font = S.LABEL_FONT
        for c in range(2, 13):
            ws_a.cell(row=ra, column=c).number_format = S.NUM_FMT
        ws_a.cell(row=ra, column=13).number_format = S.SCORE_FMT
        ws_a.cell(row=ra, column=13).alignment = S.CENTER
        ra += 1

    ra += 1
    ra = _section(ws_a, ra, 1, "Cost of Capital & Returns", span=14)
    for label, fmt in (("WACC", S.PCT_FMT), ("ROIC", S.PCT_FMT),
                       ("ROIC – WACC", S.PCT_FMT), ("ROCE", S.PCT_FMT),
                       ("CROCI", S.PCT_FMT), ("CROCI – WACC", S.PCT_FMT),
                       ("Debt/FCF", S.NUM_FMT), ("Debt/Net OCF", S.NUM_FMT),
                       ("Debt/Assets", S.PCT_FMT), ("Interest Cover", S.NUM_FMT),
                       ("FCF/Interest", S.NUM_FMT), ("Debt/OCF", S.NUM_FMT),
                       ("CAPEX % of OCF", S.PCT_FMT),
                       ("Internal Credit Rating", None),
                       ("Credit Spread", S.PCT_FMT)):
        ws_a.cell(row=ra, column=1, value=label).font = S.LABEL_FONT
        if fmt:
            for c in range(2, 13):
                ws_a.cell(row=ra, column=c).number_format = fmt
        ra += 1

    ra += 1
    ra = _section(ws_a, ra, 1, "Owner Earnings & Capital Allocation", span=14)
    for label, fmt in (("Owner Earnings (TTM)", S.USD_FMT),
                       ("Owner Earnings Margin", S.PCT_FMT),
                       ("Maintenance CAPEX (D&A proxy)", S.USD_FMT),
                       ("Growth CAPEX", S.USD_FMT),
                       ("Capital Allocation: CAPEX %", S.PCT_FMT),
                       ("Capital Allocation: Buybacks %", S.PCT_FMT),
                       ("Capital Allocation: Dividends %", S.PCT_FMT),
                       ("Capital Allocation: M&A %", S.PCT_FMT),
                       ("Capital Allocation: Debt Paydown %", S.PCT_FMT)):
        ws_a.cell(row=ra, column=1, value=label).font = S.LABEL_FONT
        ws_a.cell(row=ra, column=2).number_format = fmt
        ra += 1

    ra += 1
    ra = _section(ws_a, ra, 1, "DuPont 5-Step (Tax × Interest × OpMargin × AssetTurn × EqMult = ROE)", span=14)
    for label in ("Tax Burden", "Interest Burden", "Operating Margin",
                  "Asset Turnover", "Equity Multiplier", "Implied ROE"):
        ws_a.cell(row=ra, column=1, value=label).font = S.LABEL_FONT
        for c in range(2, 13):
            ws_a.cell(row=ra, column=c).number_format = S.PCT_FMT
        ra += 1

    ra += 1
    ra = _section(ws_a, ra, 1, "Multiples Band (current vs 3y/5y/10y)", span=8)
    headers = ["Multiple", "Current", "Median 3y", "Median 5y", "Median 10y", "Min 5y", "Max 5y", "Sector Median"]
    for i, h in enumerate(headers):
        c = ws_a.cell(row=ra, column=i + 1, value=h)
        c.fill = S.SUBSECTION_FILL
        c.font = S.LABEL_FONT
        c.alignment = S.CENTER
    ra += 1
    for m in ("P/E", "EV/EBITDA", "EV/Sales", "P/B", "FCF Yield"):
        ws_a.cell(row=ra, column=1, value=m).font = S.LABEL_FONT
        for c in range(2, 9):
            ws_a.cell(row=ra, column=c).number_format = S.NUM_FMT
        ra += 1

    # Conditional formatting on scoring column (M, col 13)
    score_range = "M5:M40"
    rule = ColorScaleRule(
        start_type="num", start_value=0, start_color="EF4444",
        mid_type="num", mid_value=5, mid_color="FBBF24",
        end_type="num", end_value=10, end_color="10B981",
    )
    ws_a.conditional_formatting.add(score_range, rule)

    # Charts section — 2x2 grid at the bottom of Analysis, referencing Fin Stat
    chart_anchor_row = ra + 2
    ws_a.cell(row=chart_anchor_row, column=1, value="Charts").font = S.SECTION_FONT
    ws_a.cell(row=chart_anchor_row, column=1).fill = S.SECTION_FILL
    ws_a.merge_cells(start_row=chart_anchor_row, start_column=1,
                     end_row=chart_anchor_row, end_column=14)

    def _quarterly_chart(chart_cls, title, fin_row, anchor_cell):
        chart = chart_cls()
        chart.title = title
        chart.height = 7.5   # cm
        chart.width = 11.5
        chart.style = 10
        # Axes must be explicitly un-deleted or openpyxl hides them.
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.x_axis.title = "Quarter"
        chart.y_axis.title = "USD"
        # min_col=1 includes the Fin Stat row label so the series is named.
        data = Reference(fws, min_col=1, max_col=13,
                         min_row=fin_row, max_row=fin_row)
        chart.add_data(data, titles_from_data=True, from_rows=True)
        cats = Reference(fws, min_col=2, max_col=13,
                         min_row=is_period_row, max_row=is_period_row)
        chart.set_categories(cats)
        if chart.legend is not None:
            chart.legend.position = "b"
        ws_a.add_chart(chart, anchor_cell)

    # 2x2 grid — left column at A, right column at J (no overlap at 11.5cm wide)
    _quarterly_chart(BarChart, "Operating Income (EBIT)",
                     is_row_of["Operating Income (EBIT)"],
                     f"A{chart_anchor_row + 1}")
    _quarterly_chart(LineChart, "Revenue",
                     is_row_of["Revenue"],
                     f"J{chart_anchor_row + 1}")
    _quarterly_chart(BarChart, "Net Income",
                     is_row_of["Net Income"],
                     f"A{chart_anchor_row + 17}")
    _quarterly_chart(LineChart, "Pretax Income",
                     is_row_of["Pretax Income"],
                     f"J{chart_anchor_row + 17}")

    # --- Valuation ---
    vws = wb.create_sheet("Valuation")
    _set_header(vws, "DCF Valuation", 14)
    _col_widths(vws, {1: 32})
    for c in range(2, 14):
        vws.column_dimensions[get_column_letter(c)].width = 12
    rv = 3
    rv = _section(vws, rv, 1, "Inputs", span=4)
    for label, fmt, inp in (("WACC", S.PCT_FMT, False),
                            ("Scenario", None, True),
                            ("Net Debt", S.USD_FMT, False),
                            ("Minority Interest", S.USD_FMT, False),
                            ("Shares Outstanding", S.INT_FMT, False),
                            ("Current Price", S.USD_FMT, False)):
        rv = _label_value(vws, rv, 1, label, fmt=fmt, input_cell=inp)

    rv += 1
    rv = _section(vws, rv, 1, "Forecast (Year +1 ... +5)", span=7)
    year_headers = ["", "Y+1", "Y+2", "Y+3", "Y+4", "Y+5", "Terminal"]
    for i, h in enumerate(year_headers):
        c = vws.cell(row=rv, column=i + 1, value=h)
        c.fill = S.SUBSECTION_FILL
        c.font = S.LABEL_FONT
        c.alignment = S.CENTER
    rv += 1
    for label, fmt in (("Revenue", S.USD_FMT), ("EBIT", S.USD_FMT),
                       ("EBITDA", S.USD_FMT), ("Tax", S.USD_FMT),
                       ("Unlevered Net Income", S.USD_FMT),
                       ("+ D&A", S.USD_FMT), ("− CAPEX", S.USD_FMT),
                       ("− Δ NWC", S.USD_FMT), ("Unlevered FCF", S.USD_FMT)):
        vws.cell(row=rv, column=1, value=label).font = S.LABEL_FONT
        for c in range(2, 8):
            vws.cell(row=rv, column=c).number_format = fmt
        rv += 1

    rv += 1
    rv = _section(vws, rv, 1, "Sensitivity: Price per Share (rows = discount rate, cols = terminal multiple)", span=6)
    mults = config.get("terminal_multiples", [8.0, 10.0, 12.0])
    rates = config.get("discount_rates", [0.03, 0.04, 0.05])
    vws.cell(row=rv, column=1, value="Disc \\ Mult").font = S.LABEL_FONT
    for i, m in enumerate(mults):
        vws.cell(row=rv, column=2 + i, value=f"{m:.1f}x").font = S.LABEL_FONT
    rv += 1
    for r_idx, rate in enumerate(rates):
        vws.cell(row=rv + r_idx, column=1, value=f"{rate*100:.0f}%").font = S.LABEL_FONT
        for c_idx in range(len(mults)):
            vws.cell(row=rv + r_idx, column=2 + c_idx).number_format = S.USD_FMT

    rv += len(rates) + 2
    rv = _section(vws, rv, 1, "Sensitivity: Implied Upside %", span=6)
    vws.cell(row=rv, column=1, value="Disc \\ Mult").font = S.LABEL_FONT
    for i, m in enumerate(mults):
        vws.cell(row=rv, column=2 + i, value=f"{m:.1f}x").font = S.LABEL_FONT
    rv += 1
    for r_idx, rate in enumerate(rates):
        vws.cell(row=rv + r_idx, column=1, value=f"{rate*100:.0f}%").font = S.LABEL_FONT
        for c_idx in range(len(mults)):
            vws.cell(row=rv + r_idx, column=2 + c_idx).number_format = S.PCT_FMT

    rv += len(rates) + 2
    rv = _section(vws, rv, 1, "Reverse DCF & Margin of Safety", span=4)
    _label_value(vws, rv, 1, "Implied Growth (current price)", fmt=S.PCT_FMT)
    _label_value(vws, rv, 4, "Margin of Safety", fmt=S.PCT_FMT)

    # --- Credit ---
    cws_c = wb.create_sheet("Credit")
    _set_header(cws_c, "Structural Credit / Default Models", 10)
    _col_widths(cws_c, {1: 36, 2: 18, 3: 6, 4: 30, 5: 18})
    rc = 3
    rc = _section(cws_c, rc, 1, "KMV — Expected Default Frequency (PRIMARY)", span=2)
    for label, fmt, inp in (("Equity Value (E)", S.USD_FMT, True),
                            ("Equity Volatility (σE)", S.PCT_FMT, True),
                            ("Risk-Free Rate", S.PCT_FMT, False),
                            ("Long-Term Liabilities", S.USD_FMT, True),
                            ("Short-Term Liabilities", S.USD_FMT, True),
                            ("Horizon (years)", S.NUM_FMT, True),
                            ("Default Point (0.5·LT + ST)", S.USD_FMT, False),
                            ("Solved Asset Value (A)", S.USD_FMT, False),
                            ("Solved Asset Volatility (σA)", S.PCT_FMT, False),
                            ("Distance to Default", S.NUM_FMT, False),
                            ("EDF (Expected Default Frequency)", S.PCT_FMT, False),
                            ("Solver Squared Error", S.NUM_FMT, False)):
        rc = _label_value(cws_c, rc, 1, label, fmt=fmt, input_cell=inp)

    rc2 = 3
    rc2 = _section(cws_c, rc2, 4, "Altman Z-Score (PRIMARY cross-check)", span=2)
    for label, fmt in (("X1 = WC / TA", S.NUM_FMT),
                       ("X2 = RE / TA", S.NUM_FMT),
                       ("X3 = EBIT / TA", S.NUM_FMT),
                       ("X4 = MVe / TL", S.NUM_FMT),
                       ("X5 = Sales / TA", S.NUM_FMT),
                       ("Z-Score", S.NUM_FMT),
                       ("Bucket", None)):
        rc2 = _label_value(cws_c, rc2, 4, label, fmt=fmt)

    rc2 += 1
    rc2 = _section(cws_c, rc2, 4, "Alternate Models (collapsed)", span=2)
    for label, fmt in (("Merton: Put Value", S.USD_FMT),
                       ("Merton: CDS per Year", S.PCT_FMT),
                       ("Merton: Default Prob (Term)", S.PCT_FMT),
                       ("Merton: Annual Default Prob", S.PCT_FMT),
                       ("Hillegeist: Score", S.NUM_FMT),
                       ("Hillegeist: Default Prob", S.PCT_FMT)):
        rc2 = _label_value(cws_c, rc2, 4, label, fmt=fmt)

    # --- Trade Notes ---
    tws = wb.create_sheet("Trade Notes")
    _set_header(tws, "Trade Notes & Decision Log", 8)
    _col_widths(tws, {1: 22, 2: 60})
    rt = 3
    for label in ("Entry Thesis", "Position Size", "Stop Level", "Time Horizon",
                  "Pre-Mortem (I'd know I'm wrong if...)", "Review Checkpoints",
                  "Insider Activity Notes", "Short Interest History"):
        tws.cell(row=rt, column=1, value=label).font = S.LABEL_FONT
        tws.merge_cells(start_row=rt, start_column=2, end_row=rt + 2, end_column=8)
        tws.cell(row=rt, column=2).alignment = S.TOP_LEFT
        tws.cell(row=rt, column=2).fill = S.INPUT_FILL
        rt += 3

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def main():
    repo_build = Path(__file__).resolve().parent.parent / "build"
    repo_build.mkdir(parents=True, exist_ok=True)
    market = build_market_daily(repo_build / "Market_Daily.xlsx")
    template = build_ticker_template(repo_build / "Ticker_TEMPLATE.xlsx")
    print(f"Built {market}")
    print(f"Built {template}")


if __name__ == "__main__":
    main()
