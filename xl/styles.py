from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


HEADER_FILL = _fill("1F2937")          # near-black
SECTION_FILL = _fill("2563EB")         # blue
SUBSECTION_FILL = _fill("DBEAFE")      # light blue
GOOD_FILL = _fill("10B981")            # green
WARN_FILL = _fill("FBBF24")            # yellow
BAD_FILL = _fill("EF4444")             # red
NEUTRAL_FILL = _fill("F3F4F6")
INPUT_FILL = _fill("FEF3C7")           # editable input cells: pale yellow

HEADER_FONT = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
LABEL_FONT = Font(name="Calibri", size=10, bold=True)
BODY_FONT = Font(name="Calibri", size=10)
MUTED_FONT = Font(name="Calibri", size=9, italic=True, color="6B7280")

THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
TOP_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)


PCT_FMT = "0.0%;[Red]-0.0%"
NUM_FMT = "#,##0.00;[Red]-#,##0.00"
INT_FMT = "#,##0;[Red]-#,##0"
SCORE_FMT = "0"
USD_FMT = "$#,##0.00;[Red]-$#,##0.00"
